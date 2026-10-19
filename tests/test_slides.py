"""Slide masters (.pptx for PowerPoint, Google Slides and the Keynote source)."""

import re
import runpy
import sys
import zipfile

import pytest

pytest.importorskip("pptx")
from PIL import Image  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER  # noqa: E402

import fwl_visual_language as fvl  # noqa: E402
from fwl_visual_language import slides  # noqa: E402

CONTENT = ("Figure", "Figure, large", "Two figures", "Team")
EMU = 12700


@pytest.fixture(scope="module")
def figure(tmp_path_factory):
    path = tmp_path_factory.mktemp("img") / "fig.png"
    Image.new("RGB", (142, 92), (255, 255, 255)).save(path)
    return path


@pytest.fixture(scope="module")
def built(tmp_path_factory, figure):
    d = tmp_path_factory.mktemp("slides")
    return {
        "pptx": slides.build_slides(d / "a.pptx", figure=figure),
        "keynote": slides.build_slides(d / "k.pptx", keynote=True, figure=figure),
        "plain": slides.build_slides(d / "p.pptx"),
    }


def layout_xml(path):
    with zipfile.ZipFile(path) as z:
        names = sorted(
            (n for n in z.namelist() if re.fullmatch(r"ppt/slideLayouts/slideLayout\d+\.xml", n)),
            key=lambda n: int(re.search(r"(\d+)\.xml", n).group(1)),
        )
        return [z.read(n).decode() for n in names]


def test_layouts_have_the_lab_names_and_sizes(built):
    for key, (w, h) in {"pptx": (960, 540), "keynote": (1920, 1080)}.items():
        prs = Presentation(built[key])
        assert [lay.name for lay in prs.slide_layouts] == list(slides.LAYOUTS)
        assert (prs.slide_width, prs.slide_height) == (w * EMU, h * EMU)
        assert len(prs.slides) == len(slides.LAYOUTS)
        assert [s.slide_layout.name for s in prs.slides] == list(slides.LAYOUTS)


def test_figure_placeholders_only_in_the_powerpoint_file(built):
    def figure_boxes(path):
        prs = Presentation(path)
        return {
            lay.name: sorted(
                ph.placeholder_format.idx
                for ph in lay.placeholders
                if ph.placeholder_format.type == PP_PLACEHOLDER.OBJECT
            )
            for lay in prs.slide_layouts
        }

    pptx = figure_boxes(built["pptx"])
    assert {k: v for k, v in pptx.items() if v} == {
        "Figure": [10],
        "Figure, large": [10],
        "Two figures": [10, 11],
    }
    assert not any(figure_boxes(built["keynote"]).values())


@pytest.mark.parametrize("key,scale", [("pptx", 0.5), ("keynote", 1.0)])
def test_reference_boxes_sit_in_both_bottom_corners_and_grow_upward(built, key, scale):
    prs = Presentation(built[key])
    for lay in prs.slide_layouts:
        phs = {ph.placeholder_format.idx: ph for ph in lay.placeholders}
        assert (20 in phs and 21 in phs) == (lay.name in CONTENT), lay.name
        if lay.name not in CONTENT:
            continue
        left, right = phs[20], phs[21]
        for ph, align in ((left, "l"), (right, "r")):
            xml = ph._element.xml
            assert 'anchor="b"' in xml  # bottom anchored: a stack grows upward
            assert f'algn="{align}"' in xml
            assert f'sz="{round(slides.REF_PT * scale * 100)}"' in xml
            assert ph.top + ph.height == round((1080 - slides.EDGE) * scale * EMU)
        assert left.left == round(slides.EDGE * scale * EMU)
        assert right.left + right.width == round((1920 - slides.EDGE) * scale * EMU)
        assert left.left + left.width < right.left  # the corners do not overlap


def test_only_lab_fonts_and_token_colours(built):
    allowed = {h[1:] for name in ("ink", "signal", "magma") for h in fvl.ramp(name).values()}
    allowed |= {h[1:] for _, h in fvl.palette("dark")} | {"FFFFFF", "000000"}
    for key in ("pptx", "keynote"):
        for xml in layout_xml(built[key]):
            assert set(re.findall(r'typeface="([^"]+)"', xml)) <= {
                "Geist",
                "Geist Light",
                "Geist Medium",
                "Geist Mono",
            }
            assert set(re.findall(r'srgbClr val="([0-9A-F]{6})"', xml)) <= allowed


def test_every_layout_has_a_plain_ink_ground_and_no_background_picture(built):
    for key in ("pptx", "keynote", "plain"):
        for name, xml in zip(slides.LAYOUTS, layout_xml(built[key]), strict=True):
            ground = xml.split("</p:bg>")[0]
            assert "<a:blipFill" not in ground, (key, name)
            step = "980" if name == "Statement" else "950"
            assert f'<a:srgbClr val="{slides.INK[step]}"/>' in ground, (key, name)


def test_letter_spacing_scales_with_the_slide(built):
    for key, spc in (("pptx", "300"), ("keynote", "600")):
        title = layout_xml(built[key])[0]
        assert f'spc="{spc}"' in title


def test_theme_has_the_lab_colour_and_font_schemes(built):
    with zipfile.ZipFile(built["pptx"]) as z:
        theme = z.read("ppt/theme/theme1.xml").decode()
    assert '<a:clrScheme name="Forming Worlds">' in theme
    accent1 = re.search(r"<a:accent1><a:srgbClr val=\"(\w+)\"", theme).group(1)
    assert accent1 == fvl.ramp("magma")["400"][1:]
    assert re.search(r'<a:majorFont><a:latin typeface="Geist Light"', theme)
    assert re.search(r'<a:minorFont><a:latin typeface="Geist"', theme)


def test_sample_figures_fit_their_boxes(built):
    for key, scale in (("pptx", 0.5), ("keynote", 1.0)):
        prs = Presentation(built[key])
        for s in prs.slides:
            pics = [sh for sh in s.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE]
            boxes = [b for (n, _), b in slides.BOX.items() if n == s.slide_layout.name]
            assert len(pics) == len(boxes), s.slide_layout.name
            for pic, (x, y, w, h) in zip(sorted(pics, key=lambda p: p.left), boxes, strict=True):
                assert x * scale * EMU - 2 <= pic.left
                assert pic.left + pic.width <= (x + w) * scale * EMU + 2
                assert y * scale * EMU - 2 <= pic.top
                assert pic.top + pic.height <= (y + h) * scale * EMU + 2
            # a sample figure replaces the empty figure box on its slide
            if boxes:
                assert not [ph for ph in s.placeholders if ph.placeholder_format.idx in (10, 11)]


def test_samples_have_no_personal_content(built):
    prs = Presentation(built["pptx"])
    text = " ".join(sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame)
    text += " ".join(
        sh.text_frame.text for lay in prs.slide_layouts for sh in lay.shapes if sh.has_text_frame
    )
    assert "SPEAKER NAME" in text and "Venue, city" in text
    assert not re.search(r"\bTim\b|TIM LICHTENBERG|Leuven", text)


REAL_NAMES = re.compile(rb"Lichtenberg|Miguel|Teske|Nicholls|Treatise", re.I)
SAMPLE_REF = re.compile(r"Author( & Author|\+) \d\d([a-z](,[a-z])?)?( \S.*)?")


@pytest.mark.parametrize("key", ["pptx", "keynote", "plain"])
def test_sample_references_are_neutral_in_every_part(built, key):
    with zipfile.ZipFile(built[key]) as z:
        names = z.namelist()
        hits = [n for n in names if n.endswith((".xml", ".rels")) and REAL_NAMES.search(z.read(n))]
    assert hits == []
    assert not [n for n in names if "notesSlide" in n or "notesMaster" in n]
    prs = Presentation(built[key])
    assert (prs.core_properties.author, prs.core_properties.last_modified_by) == (
        "Forming Worlds Lab",
        "Forming Worlds Lab",
    )
    refs = [
        para.text
        for s in prs.slides
        for sh in s.placeholders
        if sh.placeholder_format.idx in (11, 12, 13, 14, 20, 21)
        for para in sh.text_frame.paragraphs
        if para.text
    ]
    assert len(refs) == 8, refs  # 4 corner references on figure slides, 4 on the summary
    assert all(SAMPLE_REF.fullmatch(r) for r in refs), refs


def test_more_layouts_than_the_base_template_has_are_refused(tmp_path, monkeypatch):
    real = slides._layouts
    monkeypatch.setattr(slides, "_layouts", lambda *a, **k: real(*a, **k) * 2)
    with pytest.raises(RuntimeError, match="expected 20 layouts in the base template, got 11"):
        slides.build_slides(tmp_path / "x.pptx")


def test_command_line_and_module_run(tmp_path, capsys, monkeypatch, figure):
    slides.main([str(tmp_path / "a.pptx"), "--keynote", "--figure", str(figure)])
    assert capsys.readouterr().out.strip() == str(tmp_path / "a.pptx")
    assert Presentation(tmp_path / "a.pptx").slide_width == 1920 * EMU
    monkeypatch.setattr(sys, "argv", ["fwl_visual_language.slides", str(tmp_path / "b.pptx")])
    monkeypatch.delitem(sys.modules, "fwl_visual_language.slides")
    runpy.run_module("fwl_visual_language.slides", run_name="__main__")
    assert (tmp_path / "b.pptx").is_file()


def test_master_maps_text_to_the_light_colour_on_the_dark_ground(built):
    for key in ("pptx", "keynote"):
        with zipfile.ZipFile(built[key]) as z:
            master = z.read("ppt/slideMasters/slideMaster1.xml").decode()
            theme = z.read("ppt/theme/theme1.xml").decode()
        cmap = dict(re.findall(r'(\w+)="(\w+)"', re.search(r"<p:clrMap [^>]*>", master).group(0)))
        scheme = dict(re.findall(r"<a:(\w+)><a:srgbClr val=\"(\w+)\"", theme))
        text, ground = "#" + scheme[cmap["tx1"]], "#" + scheme[cmap["bg1"]]
        assert ground == fvl.ramp("ink")["950"]
        assert fvl.contrast(text, ground) > 7
        # a text box and list levels 2 and below take the master text styles: Geist on ink
        assert set(re.findall(r'typeface="([^"]+)"', master)) <= {"Geist", "Geist Light"}
        assert 'val="FFFFFF"' in master and f'val="{slides.INK["100"]}"' in master


@pytest.mark.parametrize("key,scale", [("pptx", 0.5), ("keynote", 1.0)])
def test_master_placeholders_fit_the_wide_slide(built, key, scale):
    prs = Presentation(built[key])
    width, height = prs.slide_width, prs.slide_height
    title = prs.slide_master.placeholders[0]
    assert title.width == round(1792 * scale * EMU)
    for ph in prs.slide_master.placeholders:
        assert 0 <= ph.left and ph.left + ph.width <= width, ph.name
        assert 0 <= ph.top and ph.top + ph.height <= height, ph.name


def test_document_properties_name_the_lab(built):
    props = Presentation(built["pptx"]).core_properties
    assert props.author == props.last_modified_by == "Forming Worlds Lab"
    assert props.title == "Forming Worlds Lab slide masters"
    assert props.created.year == 2026


def test_theme_accents_follow_the_dark_palette_by_name(built):
    with zipfile.ZipFile(built["pptx"]) as z:
        theme = z.read("ppt/theme/theme1.xml").decode()
    dark = dict(fvl.palette("dark"))
    for accent, name in (
        ("accent3", "fern"),
        ("accent4", "ochre"),
        ("accent5", "iris"),
        ("accent6", "garnet"),
    ):
        value = re.search(rf"<a:{accent}><a:srgbClr val=\"(\w+)\"", theme).group(1)
        assert "#" + value == dark[name], accent


def test_section_label_spacing_scales_too(built):
    for key, spc in (("pptx", "150"), ("keynote", "300")):
        section = layout_xml(built[key])[slides.LAYOUTS.index("Section")]
        assert f'spc="{spc}"' in section


def test_a_missing_python_pptx_names_the_extra(tmp_path, monkeypatch):
    monkeypatch.setitem(sys.modules, "pptx", None)
    with pytest.raises(ImportError, match=r"fwl-visual-language\[slides\]"):
        slides.build_slides(tmp_path / "x.pptx")


def test_a_bad_figure_is_refused_before_the_build(tmp_path):
    (tmp_path / "text.png").write_text("not an image")
    Image.new("RGB", (40, 30), (255, 0, 0)).save(tmp_path / "ok.png")
    data = bytearray((tmp_path / "ok.png").read_bytes())
    data[data.index(b"IDAT") + 6] ^= 0xFF  # a byte of the pixel data
    (tmp_path / "flipped.png").write_bytes(bytes(data))
    (tmp_path / "half.png").write_bytes(bytes(data[: len(data) // 2]))
    for bad in ("missing.png", "text.png", "flipped.png", "half.png"):
        with pytest.raises(ValueError, match="cannot read the figure"):
            slides.build_slides(tmp_path / "x.pptx", figure=tmp_path / bad)
    Image.new("RGB", (40, 30)).save(tmp_path / "fig.webp")
    with pytest.raises(ValueError, match="is WEBP"):
        slides.build_slides(tmp_path / "x.pptx", figure=tmp_path / "fig.webp")
    assert not (tmp_path / "x.pptx").exists()
    assert slides.build_slides(tmp_path / "x.pptx", figure=tmp_path / "ok.png").is_file()


def test_missing_output_folders_are_created(tmp_path):
    out = slides.build_slides(tmp_path / "a" / "b" / "x.pptx")
    assert out.is_file()


def test_master_body_levels_and_the_free_bottom_corners(built):
    prs = Presentation(built["keynote"])
    master = prs.slide_master.element.xml
    body = re.search(r"<p:bodyStyle>(.*?)</p:bodyStyle>", master, re.S).group(1)
    for n in range(1, 6):
        level = re.search(rf"<a:lvl{n}pPr[^>]*>(.*?)</a:lvl{n}pPr>", body, re.S).group(1)
        assert '<a:buChar char="•"/>' in level and 'typeface="Geist"' in level, n
    title_style = re.search(r"<p:titleStyle>(.*?)</p:titleStyle>", master, re.S).group(1)
    assert 'typeface="Geist Light"' in title_style
    # date and footer sit at the top, away from the reference corners
    for ph in prs.slide_master.placeholders:
        if ph.placeholder_format.type in (PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.FOOTER):
            assert ph.top + ph.height <= 40 * EMU, ph.name
    # the master and the layouts share the title and slide-number boxes
    title = prs.slide_master.placeholders[0]
    figure_title = prs.slide_layouts[slides.LAYOUTS.index("Figure")].placeholders[0]
    assert (title.left, title.top, title.width) == (
        figure_title.left,
        figure_title.top,
        figure_title.width,
    )
