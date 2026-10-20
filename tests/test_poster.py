"""Posters and one-pagers: the LaTeX style ``fwlposter.sty`` and the PowerPoint posters."""

import runpy
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path

import pytest

from fwl_visual_language import __main__ as cli
from fwl_visual_language import assets_dir, poster, ramp

TEMPLATES = Path(__file__).resolve().parents[1] / "templates"
TEX = shutil.which("lualatex") and shutil.which("pdffonts") and shutil.which("pdftotext")
MESSAGES = ["First statement of the poster", "Second statement here", "Third statement here"]


def poster_tex(fmt, look, messages=3):
    body = "".join(
        f"\\postermessage{{{MESSAGES[i % 3]}}}{{fig}}{{Ref {i}a}}{{fig}}{{Ref {i}b}}"
        for i in range(messages)
    )
    return (
        f"\\documentclass{{article}}\\usepackage[{fmt},{look}]{{fwlposter}}"
        "\\postertitle{A poster title with spaces}\\posterauthors{Name Surname and the Lab}"
        "\\posteraffiliation{An Institute}\\postercontact{formingworlds.space}"
        f"{body}\\begin{{document}}\\makeposter\\end{{document}}\n"
    )


def compile_tex(folder, name, text, ok=True):
    (folder / f"{name}.tex").write_text(text)
    for _ in range(2):  # the overlay picture reads the page position from the first run
        run = subprocess.run(
            ["lualatex", "-interaction=nonstopmode", "-halt-on-error", f"{name}.tex"],
            cwd=folder, capture_output=True, text=True, timeout=300,
        )  # fmt: skip
        if run.returncode:
            break
    assert (run.returncode == 0) == ok, run.stdout[-1500:]
    return run


def test_the_style_declares_every_format_and_both_looks():
    text = poster.poster_sty()
    for name, (w, h, base, layout, *_rest) in poster.FORMATS.items():
        assert f"\\DeclareOption{{{name}}}{{\\def\\fwl@W{{{w}}}\\def\\fwl@H{{{h}}}" in text
        assert f"\\def\\fwl@B{{{base}}}\\def\\fwl@layout{{{layout}}}" in text
    assert "\\DeclareOption{dark}" in text and "\\DeclareOption{light}" in text
    assert "\\ProvidesPackage{fwlposter}[2026/10/20" in text
    assert "\\RequirePackage[math]{fwl}[2026/10/20]" in text
    assert "text width" not in text.split("\\newcommand{\\fwl@text}")[1]


def test_install_poster_writes_the_style_fonts_and_logos(tmp_path):
    out = poster.install_poster(tmp_path / "p")
    assert out.read_text(encoding="utf-8") == poster.poster_sty()
    for rel in ("fwl.sty", "fwl-fonts/Geist-Light.ttf", "fwl-fonts/Inter-Regular.ttf"):
        assert (tmp_path / "p" / rel).is_file(), rel
    for logo in ("fwl-logo-white.png", "fwl-logo-ink.png"):
        assert (tmp_path / "p" / "fwl-assets" / logo).read_bytes() == (
            assets_dir() / logo
        ).read_bytes()


@pytest.fixture(scope="module")
def tex_folder(tmp_path_factory):
    folder = tmp_path_factory.mktemp("poster")
    poster.install_poster(folder)
    from PIL import Image

    Image.new("RGB", (800, 500), "#C0392B").save(folder / "fig.png")
    return folder


@pytest.mark.skipif(not TEX, reason="no TeX")
@pytest.mark.parametrize("fmt", list(poster.FORMATS))
@pytest.mark.parametrize("look", ["dark", "light"])
def test_each_format_compiles_to_its_page_size_with_every_text(tex_folder, fmt, look):
    name = f"{fmt}-{look}"
    run = compile_tex(tex_folder, name, poster_tex(fmt, look))
    log = (tex_folder / f"{name}.log").read_text(errors="replace")
    assert "Overfull" not in log and "Missing character" not in log, run.stdout[-500:]
    info = subprocess.run(
        ["pdfinfo", f"{name}.pdf"], cwd=tex_folder, capture_output=True, text=True
    )
    w, h = (float(v) for v in info.stdout.split("Page size:")[1].split("pts")[0].split(" x "))
    wmm, hmm = poster.FORMATS[fmt][:2]
    assert abs(w - wmm * 72 / 25.4) < 1 and abs(h - hmm * 72 / 25.4) < 1, (w, h)
    assert "Pages:           1" in info.stdout
    text = subprocess.run(
        ["pdftotext", f"{name}.pdf", "-"], cwd=tex_folder, capture_output=True, text=True
    ).stdout
    # spaces survive: a node with a set text width drops them in large type
    for words in ["A poster title with spaces", "Name Surname and the Lab", *MESSAGES, "Ref 2b"]:
        assert words in " ".join(text.split()), (words, text[:300])
    fonts = subprocess.run(
        ["pdffonts", f"{name}.pdf"], cwd=tex_folder, capture_output=True, text=True
    )
    names = {ln.split()[0].split("+")[-1] for ln in fonts.stdout.splitlines()[2:]}
    assert {"Geist-Light", "Geist-Medium", "Geist-Regular"} <= names, names
    assert not {n for n in names if not n.startswith(("Geist", "Inter", "STIX"))}, names


@pytest.mark.skipif(not TEX, reason="no TeX")
@pytest.mark.parametrize("messages", [2, 4])
def test_a_poster_needs_exactly_3_messages(tex_folder, messages):
    run = compile_tex(tex_folder, f"m{messages}", poster_tex("a4", "dark", messages), ok=False)
    assert "3 messages" in run.stdout or "needs 3" in run.stdout


@pytest.mark.skipif(not TEX, reason="no TeX")
def test_the_dark_page_is_ink_and_the_light_page_white(tex_folder):
    from PIL import Image

    for look, ground in (("dark", ramp("ink")["950"]), ("light", "#FFFFFF")):
        name = f"a4-{look}"
        if not (tex_folder / f"{name}.pdf").exists():
            compile_tex(tex_folder, name, poster_tex("a4", look))
        subprocess.run(["pdftoppm", "-r", "20", "-png", f"{name}.pdf", name], cwd=tex_folder)
        im = Image.open(next(tex_folder.glob(f"{name}-*.png"))).convert("RGB")
        corner = im.getpixel((im.width // 2, im.height - 3))
        want = tuple(int(ground[i : i + 2], 16) for i in (1, 3, 5))
        assert all(abs(a - b) <= 3 for a, b in zip(corner, want, strict=True)), (look, corner)


pptx = pytest.importorskip("pptx")


def slide_texts(slide):
    return [s.text_frame.text for s in slide.shapes if s.has_text_frame and s.text_frame.text]


@pytest.mark.parametrize("fmt", list(poster.FORMATS))
def test_the_pptx_poster_has_a_dark_and_a_light_slide_of_the_format(tmp_path, fmt):
    from pptx import Presentation
    from pptx.util import Emu

    path = poster.build_pptx_poster(tmp_path / f"{fmt}.pptx", fmt)
    prs = Presentation(path)
    wmm, hmm = poster.FORMATS[fmt][:2]
    assert (prs.slide_width, prs.slide_height) == (Emu(round(wmm * 36000)), Emu(round(hmm * 36000)))
    assert len(prs.slides) == 2
    for slide, ground in zip(prs.slides, ("010305", "FFFFFF"), strict=True):
        assert str(slide.background.fill.fore_color.rgb) == ground
        texts = slide_texts(slide)
        assert poster.EXAMPLE["title"] in texts[0]
        for statement, ref_a, _ in poster.EXAMPLE["messages"]:
            assert statement in texts and ref_a in texts
        assert {"1", "2", "3"} <= set(texts)
        pictures = [s for s in slide.shapes if s.shape_type == 13]
        assert len(pictures) == 7  # the logo and 6 figures
        # every shape lies on the slide
        for s in slide.shapes:
            assert 0 <= s.left and s.left + s.width <= prs.slide_width + 2, s.name
            assert 0 <= s.top and s.top + s.height <= prs.slide_height + 2, s.name


def test_the_pptx_poster_takes_figures_and_content(tmp_path):
    from PIL import Image
    from pptx import Presentation

    figs = []
    for k in range(6):
        figs.append(tmp_path / f"f{k}.png")
        Image.new("RGB", (300 + 100 * k, 200), "#334455").save(figs[-1])
    content = dict(poster.EXAMPLE, title="Own title")
    prs = Presentation(poster.build_pptx_poster(tmp_path / "p.pptx", "a4", figs, content))
    assert "Own title" in slide_texts(prs.slides[0])[0]
    widths = sorted({s.image.size for s in prs.slides[0].shapes if s.shape_type == 13})
    assert (300, 200) in widths and (800, 200) in widths


@pytest.mark.parametrize(
    "kwargs, match",
    [
        ({"fmt": "a3"}, "unknown format"),
        ({"figures": ["a.png"]}, "6 figures"),
        ({"content": dict(poster.EXAMPLE, messages=[("a", "b", "c")])}, "3 messages"),
    ],
)
def test_the_pptx_poster_refuses_bad_input(tmp_path, kwargs, match):
    with pytest.raises(ValueError, match=match):
        poster.build_pptx_poster(tmp_path / "p.pptx", **kwargs)


def test_the_pptx_poster_names_the_extra_without_python_pptx(tmp_path, monkeypatch):
    monkeypatch.setitem(sys.modules, "pptx", None)
    with pytest.raises(ImportError, match=r"fwl-visual-language\[slides\]"):
        poster.build_pptx_poster(tmp_path / "p.pptx")


def test_the_committed_poster_templates_match_a_fresh_build(tmp_path):
    fresh = poster.build_pptx_posters(tmp_path)
    for path in fresh:
        committed = TEMPLATES / path.name
        assert committed.is_file(), "run: python -m fwl_visual_language.poster --pptx templates"
        with zipfile.ZipFile(path) as a, zipfile.ZipFile(committed) as b:
            assert a.namelist() == b.namelist(), path.name
            for name in a.namelist():
                assert a.read(name) == b.read(name), (path.name, name)


def test_the_command_line_installs_the_style_or_writes_the_pptx(tmp_path, capsys, monkeypatch):
    poster.main([str(tmp_path / "tex")])
    assert (tmp_path / "tex" / poster.STY_NAME).is_file()
    poster.main(["--pptx", str(tmp_path / "pptx")])
    assert len(list((tmp_path / "pptx").glob("*.pptx"))) == len(poster.FORMATS)
    assert "Forming Worlds Poster (a4).pptx" in capsys.readouterr().out
    monkeypatch.setattr(sys, "argv", ["poster", str(tmp_path / "again")])
    runpy.run_module("fwl_visual_language.poster", run_name="__main__")
    assert (tmp_path / "again" / poster.STY_NAME).is_file()


def test_the_theme_build_holds_the_poster_style():
    assert cli.OUTPUTS["latex/fwlposter.sty"] is poster.poster_sty
