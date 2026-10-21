"""Word templates for letters, reports, memos and handbooks, written as ``.dotx`` files.

There are two letters: one with the lab letterhead and one with the banner of the
University of Groningen and the lab logo in the footer; both justify their text.

The handbook, for long internal documents, has a cover, a contents field and
chapters that start on a new page; the report keeps its chapters in the text
flow. Each kind comes in two faces: Geist, for documents that leave as PDF, and
Aptos, the Office default, for documents that others edit without the lab
fonts. Text, headings, lists, notes and captions use named styles (among them
Normal, Title, Subtitle, Heading 1 to 3, Caption, List Bullet, List Number,
Note, Key word, Hyperlink, Header and Footer), so text typed into a template
keeps the lab look. Pages are A4 with the ink logo for paper, ink text and
signal-blue accents.

Build them with the ``word`` extra (``python-docx``)::

    python -m fwl_visual_language.word templates

``--gdocs`` writes the Geist documents that Google Drive converts into Google
Docs, adapted to what Google Docs keeps (:func:`build_gdocs`).
"""

import argparse
from datetime import datetime
from pathlib import Path

from . import assets_dir, palette, ramp

__all__ = ["FACES", "KINDS", "build_gdocs", "build_word", "gdocs_name", "template_name"]

FACES = {
    "geist": {"regular": "Geist", "light": "Geist Light", "medium": "Geist Medium"},
    "aptos": {"regular": "Aptos", "light": "Aptos Light", "medium": "Aptos SemiBold"},
}
KINDS = ("letter", "university-letter", "report", "memo", "handbook")
TEMPLATE_CT = "application/vnd.openxmlformats-officedocument.wordprocessingml.template.main+xml"
PAGE_W, PAGE_H, SIDE = 21.0, 29.7, 2.2  # cm
TEXT_W = PAGE_W - 2 * SIDE
HEADING_NUM, BULLET_NUM, LIST_NUM = 90, 91, 92
LEFT_IN_LETTERS = (
    "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3", "Subject", "Date line",
    "Sender", "Byline", "Caption", "Header", "Footer", "TOC Heading", "TOC 1", "TOC 2",
)  # fmt: skip
LAB_LINE = "Forming Worlds Lab  ·  Kapteyn Astronomical Institute  ·  University of Groningen"
SENDER = (
    "Name Surname",
    "Role, Forming Worlds Lab",
    "Kapteyn Astronomical Institute",
    "University of Groningen",
    "Landleven 12, 9747 AD Groningen",
    "name@rug.nl  ·  formingworlds.space",
)
# pPr and rPr children that may follow the elements this module adds, in schema order
_PPR_AFTER_PBDR = (
    "w:shd",
    "w:tabs",
    "w:suppressAutoHyphens",
    "w:spacing",
    "w:ind",
    "w:jc",
    "w:rPr",
)
_RPR_AFTER_SPACING = ("w:w", "w:kern", "w:position", "w:sz", "w:szCs", "w:highlight", "w:u")


def _hex(name, step):
    return ramp(name)[step].lstrip("#")


def template_name(kind, face):
    """Return the file name of one template, for example ``Forming Worlds Letter (Geist)``.

    Parameters
    ----------
    kind : str
        One of :data:`KINDS`.
    face : str
        One of the keys of :data:`FACES`.

    Returns
    -------
    str
        The name without extension.
    """
    return f"Forming Worlds {kind.replace('-', ' ').title()} ({face.capitalize()})"


def gdocs_name(kind):
    """Return the file name of one Google Docs template, ``Forming Worlds Memo (Google Docs)``.

    Parameters
    ----------
    kind : str
        One of :data:`KINDS`.

    Returns
    -------
    str
        The name without extension.
    """
    return f"Forming Worlds {kind.replace('-', ' ').title()} (Google Docs)"


class _Builder:
    """Write one template; the python-docx imports stay inside the build call."""

    def __init__(self, face, docx, chapters="A", gdocs=False):
        self.f = FACES[face]
        self.chapters = chapters
        self.gdocs = gdocs
        self.docx = docx
        from docx.oxml import OxmlElement, parse_xml
        from docx.oxml.ns import nsdecls, qn
        from docx.shared import Cm, Pt, RGBColor

        self.E, self.parse, self.nsdecls, self.qn = OxmlElement, parse_xml, nsdecls, qn
        self.Cm, self.Pt, self.RGB = Cm, Pt, RGBColor
        self.doc = docx.Document()

    # --- low-level XML helpers -------------------------------------------------------------

    def fonts(self, rpr, name):
        """Replace the run fonts of ``rpr`` (including theme fonts) by one face."""
        old = rpr.find(self.qn("w:rFonts"))
        new = self.E("w:rFonts")
        for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            new.set(self.qn(a), name)
        if old is not None:
            old.getparent().replace(old, new)
        else:
            rpr.insert_element_before(new, *rpr_successors("w:rFonts"))

    def tracking(self, rpr, points):
        sp = self.E("w:spacing")
        sp.set(self.qn("w:val"), str(round(points * 20)))
        rpr.insert_element_before(sp, *_RPR_AFTER_SPACING)

    def border(self, ppr, side, color, eighths, space=4):
        bdr = self.E("w:pBdr")
        b = self.E(f"w:{side}")
        for k, v in (("val", "single"), ("sz", eighths), ("space", space), ("color", color)):
            b.set(self.qn(f"w:{k}"), str(v))
        bdr.append(b)
        ppr.insert_element_before(bdr, *_PPR_AFTER_PBDR)

    def right_tab(self, paragraph):
        """Clear the centre and right tabs of the Header and Footer styles; add one at the edge."""
        from docx.enum.text import WD_TAB_ALIGNMENT
        from docx.shared import Twips

        ts = paragraph.paragraph_format.tab_stops
        for twips in (4680, 9360):
            ts.add_tab_stop(Twips(twips), WD_TAB_ALIGNMENT.CLEAR)
        ts.add_tab_stop(self.Cm(TEXT_W), WD_TAB_ALIGNMENT.RIGHT)

    def leader_tab(self, style):
        """Give a contents style a right tab at the margin with a dotted leader."""
        tabs = self.parse(
            f'<w:tabs {self.nsdecls("w")}><w:tab w:val="right" w:leader="dot" '
            f'w:pos="{round(TEXT_W / 2.54 * 1440)}"/></w:tabs>'
        )
        style.element.get_or_add_pPr().insert_element_before(tabs, *_PPR_AFTER_PBDR[2:])

    def page_field(self, paragraph):
        fld = self.parse(
            f'<w:fldSimple {self.nsdecls("w")} w:instr="PAGE"><w:r><w:t>1</w:t></w:r></w:fldSimple>'
        )
        paragraph._p.append(fld)

    def no_borders(self, table):
        b = self.parse(
            f"<w:tblBorders {self.nsdecls('w')}>"
            + "".join(
                f'<w:{e} w:val="nil"/>'
                for e in ("top", "left", "bottom", "right", "insideH", "insideV")
            )
            + "</w:tblBorders>"
        )
        table._tbl.tblPr.insert_element_before(
            b, "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook", "w:tblCaption"
        )

    def full_width(self, table, first=None):
        """Fix a 2-column table at the text width; ``first`` is column 1 in cm, half by default.

        Google Docs splits a table without fixed widths into equal columns.
        """
        from docx.shared import Twips

        total = self.Cm(TEXT_W).twips
        a = self.Cm(first).twips if first else total // 2
        table.autofit = False
        tblw = table._tbl.tblPr.find(self.qn("w:tblW"))
        tblw.set(self.qn("w:type"), "dxa")
        tblw.set(self.qn("w:w"), str(total))
        for col, w in zip(table.columns, (a, total - a), strict=True):
            col.width = Twips(w)
            for cell in col.cells:
                cell.width = Twips(w)

    # --- styles ----------------------------------------------------------------------------

    def style(self, name, size=None, color=None, face=None, before=None, after=None, kind=None):
        """Return a paragraph or character style with face, size, colour and spacing set."""
        from docx.enum.style import WD_STYLE_TYPE

        styles = self.doc.styles
        if name in [s.name for s in styles]:
            st = styles[name]
        else:
            st = styles.add_style(name, kind or WD_STYLE_TYPE.PARAGRAPH)
            if st.type == WD_STYLE_TYPE.PARAGRAPH:
                st.base_style = styles["Normal"]
                st.quick_style = True
        st.font.bold = st.font.italic = None
        rpr = st.element.get_or_add_rPr()
        # Word's built-in styles carry tracking, kerning and complex-script weights of their own
        for tag in ("w:bCs", "w:iCs", "w:spacing", "w:kern", "w:szCs"):
            for el in rpr.findall(self.qn(tag)):
                rpr.remove(el)
        if size:
            st.font.size = self.Pt(size)
            szcs = self.E("w:szCs")
            szcs.set(self.qn("w:val"), str(round(size * 2)))
            rpr.insert_element_before(szcs, *rpr_successors("w:szCs"))
        if color:
            st.font.color.rgb = self.RGB.from_string(color)
        self.fonts(rpr, face or self.f["regular"])
        if before is not None:
            st.paragraph_format.space_before = self.Pt(before)
        if after is not None:
            st.paragraph_format.space_after = self.Pt(after)
        return st

    def theme(self):
        """Point the theme fonts and colours at the face and the paper palette."""
        from docx.opc.constants import RELATIONSHIP_TYPE as RT
        from lxml import etree

        part = self.doc.part.part_related_by(RT.THEME)
        root = etree.fromstring(part.blob)
        a = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
        scheme = root.find(f".//{a}fontScheme")
        scheme.set("name", "Forming Worlds Lab")
        for tag, face in (("majorFont", self.f["medium"]), ("minorFont", self.f["regular"])):
            scheme.find(f"{a}{tag}/{a}latin").set("typeface", face)
        colours = root.find(f".//{a}clrScheme")
        colours.set("name", "Forming Worlds Lab")
        for key, value in theme_colours().items():
            el = colours.find(f"{a}{key}")
            for child in list(el):
                el.remove(child)
            etree.SubElement(el, f"{a}srgbClr", val=value)
        part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)

    def numbering(self):
        """Add the heading numbers, the list bullet and the step numbers, all in signal blue."""
        num = self.doc.part.numbering_part.element
        sig, f = _hex("signal", "500"), self.f["regular"]
        big = '<w:sz w:val="40"/><w:szCs w:val="40"/>' if self.chapters == "B" else None

        def lvl(i, fmt, text, ind, style=None, font=None, left=None, size=None):
            pstyle = f'<w:pStyle w:val="{style}"/>' if style else ""
            slots = ("ascii", "hAnsi", "eastAsia", "cs")
            rfonts = (
                "<w:rFonts " + " ".join(f'w:{s}="{font}"' for s in slots) + "/>" if font else ""
            )
            return (
                f'<w:lvl w:ilvl="{i}"><w:start w:val="1"/><w:numFmt w:val="{fmt}"/>{pstyle}'
                f'<w:lvlText w:val="{text}"/><w:lvlJc w:val="left"/>'
                f'<w:pPr><w:ind w:left="{left or ind}" w:hanging="{ind}"/></w:pPr>'
                f'<w:rPr>{rfonts}<w:color w:val="{sig}"/>{size or ""}</w:rPr></w:lvl>'
            )

        heads = self.parse(
            f'<w:abstractNum {self.nsdecls("w")} w:abstractNumId="{HEADING_NUM}">'
            '<w:multiLevelType w:val="multilevel"/>'
            + lvl(0, "decimal", "%1", 567 if self.chapters == "A" else 850, "Heading1", size=big)
            + lvl(1, "decimal", "%1.%2", 567, "Heading2")
            + "</w:abstractNum>"
        )
        bullet = self.parse(
            f'<w:abstractNum {self.nsdecls("w")} w:abstractNumId="{BULLET_NUM}">'
            '<w:multiLevelType w:val="singleLevel"/>'
            + lvl(0, "bullet", "•", 284, "ListBullet", f)
            + "</w:abstractNum>"
        )
        steps = self.parse(
            f'<w:abstractNum {self.nsdecls("w")} w:abstractNumId="{LIST_NUM}">'
            '<w:multiLevelType w:val="multilevel"/>'
            + lvl(0, "decimal", "%1.", 357, "ListNumber")
            + lvl(1, "lowerLetter", "%2.", 357, "ListNumber2", left=714)
            + "</w:abstractNum>"
        )
        # Google Docs keeps no heading numbers, so its files carry no heading list
        lists = ((steps, LIST_NUM), (bullet, BULLET_NUM)) + (
            () if self.gdocs else ((heads, HEADING_NUM),)
        )
        for a, _ in lists:  # every abstractNum must come before the first num
            num.insert(0, a)
        for n in sorted(n for _, n in lists):
            num.append(
                self.parse(
                    f'<w:num {self.nsdecls("w")} w:numId="{n}">'
                    f'<w:abstractNumId w:val="{n}"/></w:num>'
                )
            )

    def link(self, style, num_id, level):
        ppr = style.element.get_or_add_pPr()
        old = ppr.find(self.qn("w:numPr"))
        if old is not None:
            ppr.remove(old)
        numpr = ppr.get_or_add_numPr()
        numpr.get_or_add_ilvl().val = level
        numpr.get_or_add_numId().val = num_id

    def styles(self):
        from docx.enum.style import WD_STYLE_TYPE

        f, ink = self.f, lambda s: _hex("ink", s)
        rpr = self.doc.styles.element.find(self.qn("w:docDefaults")).find(self.qn("w:rPrDefault"))
        self.fonts(rpr.find(self.qn("w:rPr")), f["regular"])
        # the text colour sits in the defaults, so a table style can set its own
        colour = self.E("w:color")
        colour.set(self.qn("w:val"), ink("900"))
        rpr.find(self.qn("w:rPr")).insert_element_before(colour, *rpr_successors("w:color"))
        normal = self.style("Normal", 10.5, None, after=6)
        normal.paragraph_format.line_spacing = 1.2
        for name in ("Normal", "Heading 3", "List Bullet", "List Number", "List Number 2"):
            for el in self.doc.styles[name].element.get_or_add_rPr().findall(self.qn("w:color")):
                el.getparent().remove(el)
        title = self.style("Title", 24, ink("950"), f["light"], before=18, after=2)
        subtitle = self.style("Subtitle", 12, ink("600"), after=4)
        # Google Docs styles hold no heading numbers, so Heading 1 is signal blue there instead
        h1_colour = _hex("signal", "500") if self.gdocs else ink("950")
        if self.chapters == "A":
            h1 = self.style("Heading 1", 13, h1_colour, f["medium"], before=16, after=6)
        else:
            h1 = self.style("Heading 1", 20, h1_colour, f["light"], before=0, after=18)
            h1.paragraph_format.page_break_before = True
        h2 = self.style("Heading 2", 11, ink("950"), f["medium"], before=12, after=4)
        h3 = self.style("Heading 3", 10.5, None, f["medium"], before=10, after=3)
        for st in (title, subtitle, h1, h2, h3):
            ppr = st.element.get_or_add_pPr()
            for tag in ("w:pBdr", "w:numPr"):
                old = ppr.find(self.qn(tag))
                if old is not None:
                    ppr.remove(old)
        for st in (h1, h2, h3):
            st.paragraph_format.keep_with_next = True
        self.style("Caption", 8.5, ink("600"), before=4, after=12)
        self.style("List Bullet", 10.5, None, after=3)
        for name, size, color in (("Header", 7.5, ink("500")), ("Footer", 7.5, ink("500"))):
            self.style(name, size, color, after=0)
        self.style("Byline", 9, ink("500"), after=14)
        self.style("Subject", 11, ink("950"), f["medium"], before=14, after=10)
        self.style("Date line", 9.5, ink("600"), before=18)
        self.style("Sender", 8, ink("600"), after=0)
        self.style("Memo label", 9, _hex("signal", "500"), f["medium"], after=0)
        self.style("List Number", 10.5, None, after=3)
        self.style("List Number 2", 10.5, None, after=3)
        note = self.style("Note", 10, ink("700"), before=6, after=8)
        self.border(note.element.get_or_add_pPr(), "left", _hex("signal", "500"), 18, space=8)
        note.paragraph_format.left_indent = self.Cm(0.4)
        shd = self.parse(
            f'<w:shd {self.nsdecls("w")} w:val="clear" w:color="auto" w:fill="{ink("50")}"/>'
        )
        note.element.get_or_add_pPr().insert_element_before(shd, *_PPR_AFTER_PBDR[1:])
        self.style("Cover title", 40, ink("950"), f["light"], before=110, after=4)
        self.style("Cover subtitle", 16, _hex("signal", "500"), f["light"], after=24)
        toc_head = self.style("TOC Heading", 13, ink("950"), f["medium"], before=0, after=14)
        toc_head.paragraph_format.page_break_before = False
        for name, size, color, face, indent in (
            ("TOC 1", 10.5, ink("950"), f["medium"], 0),
            ("TOC 2", 10, ink("600"), f["regular"], 0.6),
        ):
            st = self.style(name, size, color, face, before=4 if name == "TOC 1" else 0, after=2)
            st.paragraph_format.left_indent = self.Cm(indent)
            self.leader_tab(st)
        self.tracking(self.doc.styles["Memo label"].element.get_or_add_rPr(), 2)
        c = WD_STYLE_TYPE.CHARACTER
        for name, slot in (("Hyperlink", "hlink"), ("FollowedHyperlink", "folHlink")):
            self.style(name, color=theme_colours()[slot], kind=c).font.underline = True
        self.style("Key word", color=_hex("magma", "600"), face=f["medium"], kind=c)
        self.numbering()
        if not self.gdocs:
            self.link(h1, HEADING_NUM, 0)
            self.link(h2, HEADING_NUM, 1)
        self.link(self.doc.styles["List Bullet"], BULLET_NUM, 0)
        self.link(self.doc.styles["List Number"], LIST_NUM, 0)
        self.link(self.doc.styles["List Number 2"], LIST_NUM, 1)
        self.link(toc_head, 0, 0)  # numId 0: no number, though the style is based on Heading 1

    # --- page and parts --------------------------------------------------------------------

    def page(self):
        s = self.doc.sections[0]
        s.page_width, s.page_height = self.Cm(PAGE_W), self.Cm(PAGE_H)
        s.top_margin = s.bottom_margin = self.Cm(2.2)
        s.left_margin = s.right_margin = self.Cm(SIDE)
        s.header_distance = s.footer_distance = self.Cm(1.0)
        return s

    def footer(self, part, left):
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        p = part.paragraphs[0]
        p.style = self.doc.styles["Footer"]
        if self.gdocs:  # Google Docs drops the right tab: text and number sit at the right
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p.add_run(left + "  ·  ")
        else:
            self.right_tab(p)
            p.add_run(left + "\t")
        self.page_field(p)

    def logo(self, paragraph, width_cm):
        paragraph.add_run().add_picture(
            str(assets_dir() / "fwl-logo-ink.png"), width=self.Cm(width_cm)
        )

    def para(self, text="", style="Normal"):
        p = self.doc.add_paragraph(text, style=self.doc.styles[style])
        return p

    # --- the kinds -----------------------------------------------------------------------

    def letter(self):
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        s = self.page()
        s.different_first_page_header_footer = True
        hdr = s.first_page_header
        t = hdr.add_table(1, 2, self.Cm(TEXT_W))
        self.no_borders(t)
        if self.gdocs:
            self.full_width(t)
        self.logo(t.cell(0, 0).paragraphs[0], 3.6)
        cell = t.cell(0, 1)
        for i, line in enumerate(SENDER):
            p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
            p.style = self.doc.styles["Sender"]
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            r = p.add_run(line)
            if i == 0:
                r.font.name = self.f["medium"]
                r.font.color.rgb = self.RGB.from_string(_hex("ink", "900"))
        rule = hdr.paragraphs[0]
        rule._p.getparent().append(rule._p)  # the empty first paragraph goes below the table
        self.border(rule._p.get_or_add_pPr(), "bottom", _hex("signal", "500"), 8)
        for part in (s.first_page_footer, s.footer):
            self.footer(part, LAB_LINE)
        self.letter_body()

    def university_letter(self):
        """The university banner on page 1, the lab logo and the page number in the footer."""
        from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        s = self.page()
        s.different_first_page_header_footer = True
        hdr = s.first_page_header
        banner = hdr.paragraphs[0]
        banner.add_run().add_picture(str(assets_dir() / "rug-banner.png"), width=self.Cm(TEXT_W))
        sender = hdr.add_paragraph(style=self.doc.styles["Sender"])
        sender.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        sender.paragraph_format.space_before = self.Pt(10)
        r = sender.add_run(SENDER[0])
        r.font.name = self.f["medium"]
        r.font.color.rgb = self.RGB.from_string(_hex("ink", "900"))
        sender.add_run(f"  ·  {SENDER[-1]}")
        for part in (s.first_page_footer, s.footer):
            p = part.paragraphs[0]
            p.style = self.doc.styles["Footer"]
            if self.gdocs:  # Google Docs drops the right tab: logo and number in a table
                t = part.add_table(1, 2, self.Cm(TEXT_W))
                self.no_borders(t)
                self.full_width(t)
                p._p.getparent().append(p._p)  # the empty paragraph Word needs goes below
                for cell in t.rows[0].cells:
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM
                logo, num = t.cell(0, 0).paragraphs[0], t.cell(0, 1).paragraphs[0]
                logo.style = num.style = self.doc.styles["Footer"]
                num.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            else:
                self.right_tab(p)
                logo = num = p
            self.logo(logo, 2.0)
            if not self.gdocs:
                p.add_run("\t")
            self.page_field(num)
        self.letter_body()

    def letter_body(self):
        """Date, subject, salutation, text and signature; the text is justified."""
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        self.doc.styles["Normal"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        for name in LEFT_IN_LETTERS:  # styles based on Normal that must not stretch a wrapped line
            self.doc.styles[name].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        self.para("Groningen, 1 January 2027", "Date line")
        self.para("Subject of the letter", "Subject")
        self.para("Dear …,")
        self.para("Text of the letter.")
        self.para("Yours sincerely,").paragraph_format.space_before = self.Pt(8)
        p = self.para("Name Surname")
        p.paragraph_format.space_before = self.Pt(30)
        p.runs[0].font.name = self.f["medium"]
        self.para("Role, Forming Worlds Lab", "Sender")

    def report(self):
        s = self.page()
        self.logo(self.para(), 2.6)
        self.para("Title of the report", "Title")
        self.para("Subtitle", "Subtitle")
        byline = self.para("Name Surname  ·  Forming Worlds Lab  ·  Month Year", "Byline")
        self.border(byline._p.get_or_add_pPr(), "bottom", _hex("signal", "500"), 8)
        self.para("Introduction", "Heading 1")
        p = self.para("Text of the section. Mark a ")
        p.add_run("key word", self.doc.styles["Key word"])
        p.add_run(" with the Key word style and a link with the Hyperlink style.")
        self.para("Details", "Heading 2")
        self.para("A point in a list.", "List Bullet")
        self.para("Figure 1. Caption of the figure.", "Caption")
        self.footer(s.footer, "Title of the report  ·  Forming Worlds Lab")

    def memo(self):
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        s = self.page()
        t = self.doc.add_table(1, 2)
        self.no_borders(t)
        if self.gdocs:
            self.full_width(t)
        t.cell(0, 0).paragraphs[0].style = self.doc.styles["Memo label"]
        t.cell(0, 0).paragraphs[0].add_run("MEMO")
        right = t.cell(0, 1).paragraphs[0]
        right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        self.logo(right, 2.0)
        if self.gdocs:  # Google Docs joins adjacent tables into one
            self.para(style="Sender")
        rows = self.doc.add_table(4, 2)
        self.no_borders(rows)
        rows.autofit = False
        if self.gdocs:
            self.full_width(rows, 2.4)
        for col, w in zip(rows.columns, (2.4, TEXT_W - 2.4), strict=True):
            col.width = self.Cm(w)
        for i, (k, v) in enumerate(
            (
                ("To", "Lab members"),
                ("From", "Name Surname"),
                ("Date", "1 January 2027"),
                ("Subject", "Subject"),
            )
        ):
            a, b = rows.cell(i, 0).paragraphs[0], rows.cell(i, 1).paragraphs[0]
            a.style, b.style = self.doc.styles["Sender"], self.doc.styles["Normal"]
            a.add_run(k)
            r = b.add_run(v)
            if k == "Subject":
                r.font.name = self.f["medium"]
            rows.cell(i, 0).width, rows.cell(i, 1).width = self.Cm(2.4), self.Cm(TEXT_W - 2.4)
        rule = self.para()
        self.border(rule._p.get_or_add_pPr(), "bottom", _hex("ink", "300"), 4)
        self.para("Text of the memo.")
        self.para("A point in a list.", "List Bullet")
        self.footer(s.footer, "Forming Worlds Lab  ·  internal memo")

    def handbook(self):
        from docx.enum.text import WD_BREAK

        s = self.page()
        s.different_first_page_header_footer = True
        self.logo(self.para(), 5.0)
        self.para("Handbook", "Cover title")
        self.para("Forming Worlds Lab", "Cover subtitle")
        byline = self.para("Version  ·  Month Year", "Byline")
        self.border(byline._p.get_or_add_pPr(), "bottom", _hex("signal", "500"), 8)
        self.para("Opening line of the welcome:", "Subject")
        welcome = self.para("A short welcome that says what the handbook is for and how to use it.")
        welcome.add_run().add_break(WD_BREAK.PAGE)
        self.para("Contents", "TOC Heading")
        if self.gdocs:  # Google Docs makes its own contents and reads no Word field
            toc = self.para("Choose Insert > Table of contents to list the chapters here.")
        else:
            toc = self.contents_field()
        toc.add_run().add_break(WD_BREAK.PAGE)
        self.para("Chapter", "Heading 1")
        self.para("Text of the chapter.")
        self.para("Section", "Heading 2")
        self.para("A step.", "List Number")
        self.para("A sub-step.", "List Number 2")
        self.para("Subsection, for example a meeting time", "Heading 3")
        self.para("A note that stands out from the text.", "Note")
        self.para("A point in a list.", "List Bullet")
        self.footer(s.footer, "Forming Worlds Lab  ·  Handbook")

    def contents_field(self):
        """Return a paragraph with a Word contents field that lists Heading 1 and 2."""
        toc = self.para()
        toc._p.append(
            self.parse(f'<w:r {self.nsdecls("w")}><w:fldChar w:fldCharType="begin"/></w:r>')
        )
        toc._p.append(
            self.parse(
                f'<w:r {self.nsdecls("w")}><w:instrText xml:space="preserve">'
                ' TOC \\o "1-2" \\h \\z \\u </w:instrText></w:r>'
            )
        )
        toc._p.append(
            self.parse(f'<w:r {self.nsdecls("w")}><w:fldChar w:fldCharType="separate"/></w:r>')
        )
        toc.add_run("Right-click here and choose Update Field to list the chapters.")
        toc._p.append(
            self.parse(f'<w:r {self.nsdecls("w")}><w:fldChar w:fldCharType="end"/></w:r>')
        )
        return toc

    def save(self, out, template):
        props = self.doc.core_properties
        props.title, props.author, props.last_modified_by = (
            "Forming Worlds Lab template",
            "Forming Worlds Lab",
            "Forming Worlds Lab",
        )
        props.created = props.modified = datetime(2026, 10, 19)
        props.revision = 1
        props.comments = ""  # python-docx's default template names itself here
        if template:
            self.doc.part._content_type = TEMPLATE_CT
        self.doc.save(str(out))


def theme_colours():
    """Return the theme colours of the Word templates.

    Accent 1 is magma and accent 2 signal, the order of the slide theme, but at
    the paper steps; the figure palette puts signal first.

    Returns
    -------
    dict
        Theme slot to 6-digit hex without ``#``: ``dk1`` ink-900, ``lt1`` white,
        ``dk2`` ink-700, ``lt2`` ink-50, ``accent1`` to ``accent6`` the paper
        palette magma, signal, fern, ochre, iris and garnet, ``hlink`` signal-500
        and ``folHlink`` iris.
    """
    paper = {name: hexv.lstrip("#") for name, hexv in palette("paper")}
    return {
        "dk1": _hex("ink", "900"),
        "lt1": "FFFFFF",
        "dk2": _hex("ink", "700"),
        "lt2": _hex("ink", "50"),
        "accent1": paper["magma"],
        "accent2": paper["signal"],
        "accent3": paper["fern"],
        "accent4": paper["ochre"],
        "accent5": paper["iris"],
        "accent6": paper["garnet"],
        "hlink": _hex("signal", "500"),
        "folHlink": paper["iris"],
    }


def rpr_successors(tag):
    """Return the rPr child tags that come after ``tag`` in schema order."""
    order = (
        "w:rStyle", "w:rFonts", "w:b", "w:bCs", "w:i", "w:iCs", "w:caps", "w:smallCaps",
        "w:strike", "w:dstrike", "w:outline", "w:shadow", "w:emboss", "w:imprint",
        "w:noProof", "w:snapToGrid", "w:vanish", "w:webHidden", "w:color", "w:spacing",
        "w:w", "w:kern", "w:position", "w:sz", "w:szCs", "w:highlight", "w:u",
    )  # fmt: skip
    return order[order.index(tag) + 1 :]


def _docx():
    try:
        import docx
    except ImportError as err:
        raise ImportError(
            'the Word templates need python-docx: pip install "fwl-visual-language[word]"'
        ) from err
    return docx


def _build(face, kind, docx, gdocs=False):
    b = _Builder(face, docx, "B" if kind == "handbook" else "A", gdocs)
    b.styles()
    b.theme()
    getattr(b, kind.replace("-", "_"))()
    return b


def build_word(folder, faces=tuple(FACES), kinds=KINDS, template=True):
    """Write the Word templates.

    Parameters
    ----------
    folder : str or pathlib.Path
        Output folder; it is created when missing.
    faces : iterable of str, optional
        Keys of :data:`FACES`; both by default.
    kinds : iterable of str, optional
        Members of :data:`KINDS`; all by default.
    template : bool, optional
        Write ``.dotx`` templates; ``False`` writes the same content as
        ``.docx`` documents, which Word opens for editing (used for renders).

    Returns
    -------
    list of pathlib.Path
        The written files, faces in the outer loop.

    Raises
    ------
    ImportError
        When ``python-docx`` is missing.
    ValueError
        When a face or kind is unknown.
    """
    docx = _docx()
    faces, kinds = list(faces), list(kinds)
    bad = [f for f in faces if f not in FACES] + [k for k in kinds if k not in KINDS]
    if bad:
        raise ValueError(f"unknown face or kind: {bad}; faces {list(FACES)}, kinds {list(KINDS)}")
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    written = []
    for face in faces:
        for kind in kinds:
            b = _build(face, kind, docx)
            out = folder / f"{template_name(kind, face)}.{'dotx' if template else 'docx'}"
            b.save(out, template)
            written.append(out)
    return written


def build_gdocs(folder, kinds=KINDS):
    """Write the Geist documents that Google Drive converts into Google Docs templates.

    They differ from the Word templates where Google Docs differs from Word:
    headings carry no numbers and Heading 1 is signal blue, footers hold no tab
    stops (text and page number sit at the right; the university letter puts its
    logo and number in a table), the memo tables have fixed widths, and the
    handbook asks for Google's own table of contents instead of a Word field.

    Parameters
    ----------
    folder : str or pathlib.Path
        Output folder; it is created when missing.
    kinds : iterable of str, optional
        Members of :data:`KINDS`; all by default.

    Returns
    -------
    list of pathlib.Path
        The written ``.docx`` files.

    Raises
    ------
    ImportError
        When ``python-docx`` is missing.
    ValueError
        When a kind is unknown.
    """
    docx = _docx()
    kinds = list(kinds)
    bad = [k for k in kinds if k not in KINDS]
    if bad:
        raise ValueError(f"unknown kind: {bad}; kinds {list(KINDS)}")
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    written = []
    for kind in kinds:
        out = folder / f"{gdocs_name(kind)}.docx"
        _build("geist", kind, docx, gdocs=True).save(out, False)
        written.append(out)
    return written


def main(argv=None):
    """Command line entry point: ``python -m fwl_visual_language.word FOLDER``."""
    parser = argparse.ArgumentParser(
        prog="python -m fwl_visual_language.word",
        description="Write the Forming Worlds Lab Word templates (.dotx), or with --gdocs "
        "the Geist .docx files for Google Docs.",
    )
    parser.add_argument("folder", help="output folder")
    parser.add_argument("--face", choices=[*FACES, "all"], help="face of the Word templates")
    parser.add_argument("--docx", action="store_true", help="write .docx documents, not templates")
    parser.add_argument(
        "--gdocs", action="store_true", help="write the Geist .docx files for Google Docs"
    )
    args = parser.parse_args(argv)
    if args.gdocs and (args.face or args.docx):
        parser.error("--gdocs writes Geist .docx files and takes neither --face nor --docx")
    faces = list(FACES) if args.face in (None, "all") else [args.face]
    if args.gdocs:
        written = build_gdocs(args.folder)
    else:
        written = build_word(args.folder, faces, template=not args.docx)
    for path in written:
        print(path)


if __name__ == "__main__":
    main()
