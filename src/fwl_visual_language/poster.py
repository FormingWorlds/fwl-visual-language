"""LaTeX poster style ``fwlposter.sty``: lab posters and one-pagers from 3 messages.

A poster carries a header with the title, the authors and the lab logo, then 3
numbered messages, each a statement with 2 figures and their references, and a
contact line at the foot. The format is an option: ``a0portrait`` (default),
``a0landscape``, ``a4`` (a one-pager) or ``slide`` (16:9, for e-poster screens);
the look is ``dark`` (default, the ink ground) or ``light`` (white). Portrait
formats stack the messages, landscape formats set them side by side. Compile
with LuaLaTeX::

    \\documentclass{article}
    \\usepackage[a0portrait,dark]{fwlposter}
    \\postertitle{Title as one statement}
    \\posterauthors{Name Surname and the Forming Worlds Lab}
    \\posteraffiliation{Kapteyn Astronomical Institute, University of Groningen}
    \\postercontact{formingworlds.space}
    \\postermessage{Statement}{figure-a}{Reference}{figure-b}{Reference}
    ... 2 more \\postermessage ...
    \\begin{document}\\makeposter\\end{document}

Install it next to a poster::

    python -m fwl_visual_language.poster path/to/poster
"""

import argparse
import shutil
from datetime import date, datetime
from pathlib import Path

from . import __version__, assets_dir, fonts_dir
from .beamer import ASSET_DIR
from .latex import FONT_DIR, install_latex

__all__ = [
    "EXAMPLE",
    "FORMATS",
    "STY_NAME",
    "build_pptx_poster",
    "build_pptx_posters",
    "install_poster",
    "poster_sty",
]

STY_NAME = "fwlposter.sty"
# name: (width mm, height mm, design units across, layout, header, title, statement, reference)
FORMATS = {
    "a0portrait": (841, 1189, 1000, "rows", 300, 50, 31, 15),
    "a0landscape": (1189, 841, 1414, "columns", 200, 46, 28, 14),
    "a4": (210, 297, 1000, "rows", 190, 40, 27, 14),
    "slide": (338.67, 190.5, 1600, "columns", 170, 40, 25, 13),
}
_LOGOS = ("fwl-logo-white.png", "fwl-logo-ink.png")
_LIGHT = ("Geist-Light.ttf", "Geist-LightItalic.ttf")

_BODY = r"""\@ifpackageloaded{fwl}{\@ifpackagelater{fwl}{2026/10/20}{}%
  {\PackageError{fwlposter}{fwl.sty is older than this style}{Install both again.}}%
  \fwlusemath}{\RequirePackage[math]{fwl}[2026/10/20]}
\RequirePackage{tikz}
\RequirePackage{graphicx}
\providecommand{\fwlassetpath}{fwl-assets/}
\RequirePackage[paperwidth=\fwl@W mm,paperheight=\fwl@H mm,margin=0pt]{geometry}
\pagestyle{empty}
\setlength{\parindent}{0pt}
\newfontfamily\fwllightface{Geist}[Path=\fwlfontpath,UprightFont=Geist-Light.ttf,
  ItalicFont=Geist-LightItalic.ttf]

% Design units: the page is \fwl@B units wide; \fwl@u gives a length, \fwl@fs a font size.
\newcommand{\fwl@u}[1]{\fpeval{(#1)*\fwl@W/\fwl@B}mm}
\newcommand{\fwl@fs}[2][1.15]{\fontsize{\fpeval{(#2)*\fwl@W/\fwl@B*2.84528}pt}%
  {\fpeval{(#2)*\fwl@W/\fwl@B*2.84528*#1}pt}\selectfont}
\iffwl@dark
  \colorlet{fwlPosterGround}{fwlInk950}\colorlet{fwlPosterHead}{fwlInk900}
  \colorlet{fwlPosterRule}{fwlSignal400}\colorlet{fwlPosterText}{white}
  \colorlet{fwlPosterMuted}{fwlInk300}\colorlet{fwlPosterLine}{fwlInk700}
  \def\fwl@logo{fwl-logo-white.png}\def\fwl@pad{8}\def\fwl@panel{white}
\else
  \colorlet{fwlPosterGround}{white}\colorlet{fwlPosterHead}{fwlInk50}
  \colorlet{fwlPosterRule}{fwlSignal500}\colorlet{fwlPosterText}{fwlInk950}
  \colorlet{fwlPosterMuted}{fwlInk600}\colorlet{fwlPosterLine}{fwlInk200}
  \def\fwl@logo{fwl-logo-ink.png}\def\fwl@pad{0}\def\fwl@panel{fwlPosterGround}
\fi
\colorlet{fwlPosterNumber}{fwlPosterRule}
\pagecolor{fwlPosterGround}

\newcommand{\postertitle}[1]{\def\fwl@title{#1}}
\newcommand{\posterauthors}[1]{\def\fwl@authors{#1}}
\newcommand{\posteraffiliation}[1]{\def\fwl@affiliation{#1}}
\newcommand{\postercontact}[1]{\def\fwl@contact{#1}}
\postertitle{}\posterauthors{}\posteraffiliation{}\postercontact{}
\newcount\fwl@messages
% \postermessage{statement}{figure}{reference}{figure}{reference}
\newcommand{\postermessage}[5]{\global\advance\fwl@messages by 1
  \ifnum\fwl@messages>3 \PackageError{fwlposter}{A poster holds 3 messages}{Remove one.}\fi
  \expandafter\def\csname fwl@m\the\fwl@messages\endcsname{#1}%
  \expandafter\def\csname fwl@fa\the\fwl@messages\endcsname{#2}%
  \expandafter\def\csname fwl@ra\the\fwl@messages\endcsname{#3}%
  \expandafter\def\csname fwl@fb\the\fwl@messages\endcsname{#4}%
  \expandafter\def\csname fwl@rb\the\fwl@messages\endcsname{#5}}

% Text set in a \parbox: a TikZ node with "text width" drops the spaces of text above 30 pt.
\newcommand{\fwl@text}[2]{\parbox[t]{\fwl@u{#1}}{\raggedright #2}}

% A figure fitted into a box (x, y top-left, width, height in design units),
% on a white panel on the dark ground, with its reference under it.
\newcommand{\fwl@figure}[6]{%
  \node[anchor=center,inner sep=\fwl@u{\fwl@pad},fill=\fwl@panel,rounded corners=\fwl@u{8}]
    at (\fwl@u{#1+#3/2},-\fwl@u{#2+#4/2})
    {\includegraphics[width=\fwl@u{#3-16},height=\fwl@u{#4-16},keepaspectratio]{#5}};
  \node[anchor=north west,inner sep=0pt,text=fwlPosterMuted] at (\fwl@u{#1},-\fwl@u{#2+#4+6})
    {\fwl@text{#3}{\fwl@fs{\fwl@ref}#6}};}

% Statements are set before the picture, where TikZ has no font; their heights go
% into \fwl@sh1 to \fwl@sh3 in design units.
\expandafter\newsavebox\csname fwl@sb1\endcsname
\expandafter\newsavebox\csname fwl@sb2\endcsname
\expandafter\newsavebox\csname fwl@sb3\endcsname
\newcommand{\fwl@measure}[1]{%
  \expandafter\sbox\csname fwl@sb#1\endcsname{\fwl@text{\fwl@sw}{\fwlmedium
    \fwl@fs[1.12]{\fwl@stmt}\color{fwlPosterText}\csname fwl@m#1\endcsname}}%
  \expandafter\xdef\csname fwl@sh#1\endcsname{\fpeval{max((\the\ht\csname fwl@sb#1\endcsname
    +\the\dp\csname fwl@sb#1\endcsname)/2.84528*\fwl@B/\fwl@W,\fwl@stmt*1.8)}}}

\newcommand{\makeposter}{%
  \ifnum\fwl@messages=3\else\PackageError{fwlposter}{A poster needs 3 \string\postermessage}{}\fi
  \edef\fwl@colw{\fpeval{(\fwl@B-80-2*34)/3}}%
  \edef\fwl@sw{\fpeval{\ifx\fwl@layout\fwl@rowsname \fwl@B-80\else\fwl@colw\fi-\fwl@stmt*1.6}}%
  \fwl@measure{1}\fwl@measure{2}\fwl@measure{3}%
  \begin{tikzpicture}[remember picture,overlay,shift={(current page.north west)}]
  \fill[fwlPosterHead] (0,0) rectangle (\paperwidth,-\fwl@u{\fwl@head});
  \fill[fwlPosterRule] (0,-\fwl@u{\fwl@head-6}) rectangle (\paperwidth,-\fwl@u{\fwl@head});
  \node[anchor=north east,inner sep=0pt] at (\fwl@u{\fwl@B-40},-\fwl@u{\fwl@head*0.2})
    {\includegraphics[height=\fwl@u{\fwl@head*0.36}]{\fwlassetpath\fwl@logo}};
  \node[anchor=north west,inner sep=0pt,text=fwlPosterText] at (\fwl@u{40},-\fwl@u{\fwl@head*0.15})
    {\fwl@text{\fwl@B*0.74-40}{%
    {\fwllightface\fwl@fs[1.08]{\fwl@titlesize}\fwl@title\par}\vspace{\fwl@u{12}}%
    {\fwlmedium\fwl@fs{\fwl@titlesize*0.44}\fwl@authors\par}\vspace{\fwl@u{4}}%
    {\color{fwlPosterMuted}\fwl@fs{\fwl@titlesize*0.36}\fwl@affiliation\par}}};
  \csname fwl@layout@\fwl@layout\endcsname
  \node[anchor=south west,inner sep=0pt,text=fwlPosterMuted]
    at (\fwl@u{40},-\fwl@u{\fwl@H/\fwl@W*\fwl@B-24}) {\fwl@fs{\fwl@ref*1.15}\fwl@contact};
  \end{tikzpicture}\null}
\def\fwl@rowsname{rows}

% The number and the statement of message #1 at x = #2, y = \fwl@y.
\newcommand{\fwl@statement}[2]{%
  \edef\fwl@sh{\csname fwl@sh#1\endcsname}%
  \node[anchor=north west,inner sep=0pt,text=fwlPosterNumber] at (\fwl@u{#2},-\fwl@u{\fwl@y})
    {\fwllightface\fwl@fs[1]{\fwl@stmt*1.9}#1};
  \node[anchor=north west,inner sep=0pt]
    at (\fwl@u{#2+\fwl@stmt*1.6},-\fwl@u{\fwl@y+\fwl@stmt*0.25})
    {\usebox{\csname fwl@sb#1\endcsname}};}

\newcommand{\fwl@layout@rows}{%
  \edef\fwl@sec{\fpeval{(\fwl@H/\fwl@W*\fwl@B-\fwl@head-36-110)/3}}%
  \edef\fwl@half{\fpeval{(\fwl@B-80-24)/2}}%
  \foreach \i in {1,2,3}{%
    \edef\fwl@y{\fpeval{\fwl@head+36+(\i-1)*\fwl@sec}}%
    \fwl@statement{\i}{40}%
    \edef\fwl@fy{\fpeval{\fwl@y+\fwl@sh+18}}%
    \edef\fwl@fh{\fpeval{\fwl@y+\fwl@sec-\fwl@ref*3.6-\fwl@fy}}%
    \fwl@figure{40}{\fwl@fy}{\fwl@half}{\fwl@fh}{\csname fwl@fa\i\endcsname}%
      {\csname fwl@ra\i\endcsname}%
    \fwl@figure{40+\fwl@half+24}{\fwl@fy}{\fwl@half}{\fwl@fh}{\csname fwl@fb\i\endcsname}%
      {\csname fwl@rb\i\endcsname}%
    \ifnum\i<3 \fill[fwlPosterLine] (\fwl@u{40},-\fwl@u{\fwl@y+\fwl@sec-12})
      rectangle (\fwl@u{\fwl@B-40},-\fwl@u{\fwl@y+\fwl@sec-11});\fi}}

\newcommand{\fwl@layout@columns}{%
  \edef\fwl@y{\fpeval{\fwl@head+30}}%
  \edef\fwl@bottom{\fpeval{\fwl@H/\fwl@W*\fwl@B-70}}%
  \foreach \i in {1,2,3}{%
    \edef\fwl@x{\fpeval{40+(\i-1)*(\fwl@colw+34)}}%
    \fwl@statement{\i}{\fwl@x}%
    \edef\fwl@fy{\fpeval{\fwl@y+max(\fwl@sh,\fwl@stmt*3.4)+16}}%
    \edef\fwl@fh{\fpeval{(\fwl@bottom-\fwl@fy-2*\fwl@ref*2.4)/2}}%
    \fwl@figure{\fwl@x}{\fwl@fy}{\fwl@colw}{\fwl@fh}{\csname fwl@fa\i\endcsname}%
      {\csname fwl@ra\i\endcsname}%
    \fwl@figure{\fwl@x}{\fwl@fy+\fwl@fh+\fwl@ref*2.4}{\fwl@colw}{\fwl@fh}%
      {\csname fwl@fb\i\endcsname}{\csname fwl@rb\i\endcsname}%
    \ifnum\i<3 \fill[fwlPosterLine] (\fwl@u{\fwl@x+\fwl@colw+16},-\fwl@u{\fwl@y})
      rectangle (\fwl@u{\fwl@x+\fwl@colw+17},-\fwl@u{\fwl@bottom});\fi}}
\endinput
"""


def poster_sty(release=None):
    """Return the text of ``fwlposter.sty``.

    Parameters
    ----------
    release : datetime.date, optional
        Date written into ``\\ProvidesPackage``. Defaults to 2026-10-20 so the
        output does not change from day to day.

    Returns
    -------
    str
        LaTeX package source with the format options of :data:`FORMATS` and the
        ``dark`` and ``light`` looks.
    """
    release = release or date(2026, 10, 20)
    head = [
        "% Forming Worlds Lab poster style (fwl-visual-language "
        f"{__version__}). Generated; edit the generator, not this file.",
        "\\NeedsTeXFormat{LaTeX2e}",
        f"\\ProvidesPackage{{fwlposter}}[{release:%Y/%m/%d} v{__version__} Forming Worlds Lab]",
        "\\newif\\iffwl@dark\\fwl@darktrue",
        "\\DeclareOption{dark}{\\fwl@darktrue}",
        "\\DeclareOption{light}{\\fwl@darkfalse}",
    ]
    for name, (w, h, base, layout, header, title, stmt, ref) in FORMATS.items():
        head.append(
            f"\\DeclareOption{{{name}}}{{\\def\\fwl@W{{{w}}}\\def\\fwl@H{{{h}}}"
            f"\\def\\fwl@B{{{base}}}\\def\\fwl@layout{{{layout}}}\\def\\fwl@head{{{header}}}"
            f"\\def\\fwl@titlesize{{{title}}}\\def\\fwl@stmt{{{stmt}}}\\def\\fwl@ref{{{ref}}}}}"
        )
    head += ["\\ExecuteOptions{a0portrait,dark}", "\\ProcessOptions\\relax", ""]
    return "\n".join(head) + _BODY


def install_poster(dest):
    """Write ``fwlposter.sty``, the ``fwl`` package, the fonts and the logos into ``dest``.

    Parameters
    ----------
    dest : str or pathlib.Path
        Folder of the poster. It is created if it does not exist.

    Returns
    -------
    pathlib.Path
        Path of the written style file.
    """
    dest = Path(dest)
    install_latex(dest)
    for face in _LIGHT:
        shutil.copyfile(fonts_dir() / "ttf" / face, dest / FONT_DIR / face)
    assets = dest / ASSET_DIR
    assets.mkdir(exist_ok=True)
    for logo in _LOGOS:
        shutil.copyfile(assets_dir() / logo, assets / logo)
    out = dest / STY_NAME
    out.write_text(poster_sty(), encoding="utf-8")
    return out


EXAMPLE = {
    "title": "Title of the poster as one statement",
    "authors": "Name Surname and the Forming Worlds Lab",
    "affiliation": "Kapteyn Astronomical Institute, University of Groningen",
    "contact": "formingworlds.space",
    "messages": [
        ("First message as one statement", "Author+ YY Journal", "Author+ YY Journal"),
        ("Second message as one statement", "Author+ YY Journal", "Author+ YY Journal"),
        ("Third message as one statement", "Author+ YY Journal", "Author+ YY Journal"),
    ],
}
_LOOKS = {
    # ground, header, rule, text, muted, line, logo, figure panel
    "dark": ("950", "900", ("signal", "400"), "0", "300", "700", "fwl-logo-white.png", True),
    "light": ("0", "50", ("signal", "500"), "950", "600", "200", "fwl-logo-ink.png", False),
}


def _placeholder():
    """Return a neutral 16:10 figure placeholder as an in-memory PNG."""
    import io

    from PIL import Image, ImageDraw

    from . import ramp

    im = Image.new("RGB", (1600, 1000), ramp("ink")["100"])
    draw = ImageDraw.Draw(im)
    for k in range(0, 1600, 80):
        draw.line((k, 1000, k + 400, 0), fill=ramp("ink")["200"], width=4)
    data = io.BytesIO()
    im.save(data, format="PNG")
    return data


def build_pptx_poster(out, fmt="a0portrait", figures=None, content=None):
    """Write a PowerPoint poster with a dark and a light slide in the given format.

    Every text is a text box and every figure a picture, so a poster is made by
    editing the slide of the chosen look and deleting the other.

    Parameters
    ----------
    out : str or pathlib.Path
        Path of the ``.pptx`` to write; missing parent folders are created.
    fmt : str
        A key of :data:`FORMATS`.
    figures : list of str or pathlib.Path, optional
        6 images in message order (2 per message); the sample figure of the
        slide templates by default.
    content : dict, optional
        Texts with the keys of :data:`EXAMPLE`; :data:`EXAMPLE` by default.

    Returns
    -------
    pathlib.Path
        Path of the written file.

    Raises
    ------
    ImportError
        If ``python-pptx`` is not installed; the message names the ``slides`` extra.
    ValueError
        If ``fmt`` is unknown, or ``figures`` does not hold 6 images, or the
        content does not hold 3 messages.
    """
    try:
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.enum.text import MSO_ANCHOR
        from pptx.util import Emu, Pt
    except ImportError as err:
        raise ImportError(
            'the posters need python-pptx: pip install "fwl-visual-language[slides]"'
        ) from err
    from PIL import Image

    from . import ramp

    if fmt not in FORMATS:
        raise ValueError(f"unknown format {fmt!r}; formats: {', '.join(FORMATS)}")
    content = content or EXAMPLE
    if len(content["messages"]) != 3:
        raise ValueError("a poster holds 3 messages")
    figures = [Path(f) for f in figures] if figures else [_placeholder()] * 6
    if len(figures) != 6:
        raise ValueError("a poster takes 6 figures, 2 per message")
    w_mm, h_mm, base, layout, header, title_u, stmt_u, ref_u = FORMATS[fmt]
    mm = w_mm / base  # mm per design unit
    emu = lambda v: Emu(round(v * mm * 36000))  # noqa: E731
    pt = lambda v: Pt(round(v * mm * 2.84528, 1))  # noqa: E731
    height_u = h_mm / mm
    ink = ramp("ink")
    rgb = lambda hexv: RGBColor.from_string(hexv.lstrip("#"))  # noqa: E731

    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(round(w_mm * 36000)), Emu(round(h_mm * 36000))
    props = prs.core_properties
    props.title = f"Forming Worlds Lab poster ({fmt})"
    props.author = props.last_modified_by = "Forming Worlds Lab"
    props.created = props.modified = datetime(2026, 10, 20)
    props.revision = 1
    blank = prs.slide_layouts[6]

    def text(slide, x, y, w, h, runs, anchor=MSO_ANCHOR.TOP):
        box = slide.shapes.add_textbox(emu(x), emu(y), emu(w), emu(h))
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        for k, (words, face, size, colour) in enumerate(runs):
            para = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            run = para.add_run()
            run.text = words
            run.font.name, run.font.size, run.font.color.rgb = face, pt(size), rgb(colour)
        return box

    def rect(slide, x, y, w, h, colour, rounded=False):
        shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
            emu(x),
            emu(y),
            emu(w),
            emu(h),
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(colour)
        shape.line.fill.background()
        shape.shadow.inherit = False
        if rounded:
            shape.adjustments[0] = min(0.5, 8 / min(w, h))
        return shape

    def figure(slide, path, x, y, w, h, panel):
        if hasattr(path, "seek"):
            path.seek(0)
        with Image.open(path) as im:
            aspect = im.width / im.height
        if hasattr(path, "seek"):
            path.seek(0)
        pad = 8 if panel else 0
        iw, ih = w - 2 * pad, h - 2 * pad
        fw, fh = (iw, iw / aspect) if iw / ih < aspect else (ih * aspect, ih)
        fx, fy = x + (w - fw) / 2, y + (h - fh) / 2
        if panel:
            rect(slide, fx - pad, fy - pad, fw + 2 * pad, fh + 2 * pad, "#FFFFFF", rounded=True)
        slide.shapes.add_picture(
            path if hasattr(path, "seek") else str(path), emu(fx), emu(fy), emu(fw), emu(fh)
        )

    for ground, head, rule, fg, muted, line, logo, panel in _LOOKS.values():
        slide = prs.slides.add_slide(blank)
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = rgb(ink[ground])
        rect(slide, 0, 0, base, header, ink[head])
        rect(slide, 0, header - 6, base, 6, ramp(rule[0])[rule[1]])
        logo_path = assets_dir() / logo
        with Image.open(logo_path) as im:
            lw = header * 0.36 * im.width / im.height
        slide.shapes.add_picture(
            str(logo_path), emu(base - 40 - lw), emu(header * 0.2), emu(lw), emu(header * 0.36)
        )
        text(
            slide,
            40,
            header * 0.15,
            base * 0.74 - 80,
            header * 0.8,
            [
                (content["title"], "Geist Light", title_u, ink[fg]),
                (content["authors"], "Geist Medium", title_u * 0.44, ink[fg]),
                (content["affiliation"], "Geist", title_u * 0.36, ink[muted]),
            ],
        )
        number = ramp(rule[0])[rule[1]]
        figs = iter(figures)
        colw = (base - 80 - 2 * 34) / 3
        for i, (statement, ref_a, ref_b) in enumerate(content["messages"]):
            if layout == "rows":
                sec = (height_u - header - 36 - 110) / 3
                y, x, width = header + 36 + i * sec, 40, base - 80
                fy = y + stmt_u * 3.4
                fh = y + sec - ref_u * 3.6 - fy
                half = (base - 80 - 24) / 2
                boxes = [(40, fy, half, fh), (40 + half + 24, fy, half, fh)]
                if i < 2:
                    rect(slide, 40, y + sec - 12, base - 80, 1, ink[line])
            else:
                y, x, width = header + 30, 40 + i * (colw + 34), colw
                bottom = height_u - 70
                fy = y + stmt_u * 3.4 + 16
                fh = (bottom - fy - 2 * ref_u * 2.4) / 2
                boxes = [(x, fy, colw, fh), (x, fy + fh + ref_u * 2.4, colw, fh)]
                if i < 2:
                    rect(slide, x + colw + 16, y, 1, bottom - y, ink[line])
            text(
                slide,
                x,
                y,
                stmt_u * 1.6,
                stmt_u * 2,
                [(str(i + 1), "Geist Light", stmt_u * 1.9, number)],
            )
            text(
                slide,
                x + stmt_u * 1.6,
                y + stmt_u * 0.25,
                width - stmt_u * 1.6,
                stmt_u * 2.6,
                [(statement, "Geist Medium", stmt_u, ink[fg])],
            )
            for (bx, by, bw, bh), ref in zip(boxes, (ref_a, ref_b), strict=True):
                figure(slide, next(figs), bx, by, bw, bh, panel)
                text(slide, bx, by + bh + 6, bw, ref_u * 2, [(ref, "Geist", ref_u, ink[muted])])
        text(
            slide,
            40,
            height_u - 24 - ref_u * 1.6,
            base - 80,
            ref_u * 1.6,
            [(content["contact"], "Geist", ref_u * 1.15, ink[muted])],
            anchor=MSO_ANCHOR.BOTTOM,
        )
    out = Path(out)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    return out


def build_pptx_posters(folder):
    """Write the PowerPoint posters of every format into ``folder``.

    Parameters
    ----------
    folder : str or pathlib.Path
        Output folder; it is created if it does not exist.

    Returns
    -------
    list of pathlib.Path
        The written files, named ``Forming Worlds Poster (<format>).pptx``.
    """
    return [
        build_pptx_poster(Path(folder) / f"Forming Worlds Poster ({fmt}).pptx", fmt)
        for fmt in FORMATS
    ]


def main(argv=None):
    """Command line: ``python -m fwl_visual_language.poster DEST`` or ``--pptx DEST``."""
    parser = argparse.ArgumentParser(
        prog="python -m fwl_visual_language.poster",
        description="Write the Forming Worlds Lab poster style with its fonts and logos, "
        "or with --pptx the PowerPoint posters.",
    )
    parser.add_argument("dest", help="folder of the poster")
    parser.add_argument("--pptx", action="store_true", help="write the PowerPoint posters")
    args = parser.parse_args(argv)
    for path in build_pptx_posters(args.dest) if args.pptx else [install_poster(args.dest)]:
        print(path)


if __name__ == "__main__":
    main()
