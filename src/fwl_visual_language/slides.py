"""Slide masters for PowerPoint, Google Slides and Keynote, written as a ``.pptx``.

The file holds 10 layouts and one sample slide per layout: Title, Question,
Section, Figure, Figure large, Two figures, Summary, Statement, Team and
Closing, all on the ink-950 ground (the statement on ink-980), with no
background picture: a talk adds its own. Every content layout has two
reference boxes at the bottom corners, 24 pt on the 1920 x 1080 design,
anchored at the bottom so a stack grows upward.

The PowerPoint and Google Slides file is 13.333 x 7.5 in with content
placeholders for figures, not picture placeholders. The Keynote source
(``keynote=True``) is 1920 x 1080 pt without figure placeholders, which
Keynote removes on import with a warning; Keynote turns the layouts into
master slides, and a saved copy with the ``.kth`` extension is the Keynote
theme.

Build it with the ``slides`` extra (``python-pptx``)::

    python -m fwl_visual_language.slides "Forming Worlds.pptx"
"""

import argparse
from datetime import datetime
from pathlib import Path
from xml.sax.saxutils import escape

from . import assets_dir, palette, ramp

__all__ = ["LAYOUTS", "build_slides"]

LAYOUTS = (
    "Title",
    "Question",
    "Section",
    "Figure",
    "Figure, large",
    "Two figures",
    "Summary",
    "Statement",
    "Team",
    "Closing",
)
_NS = (
    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
    'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
    'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"'
)
_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
INK = {step: hexv[1:] for step, hexv in ramp("ink").items()}
SIGNAL400 = ramp("signal")["400"][1:]
SIGNAL300 = ramp("signal")["300"][1:]
MAGMA400 = ramp("magma")["400"][1:]
DARK = {name: hexv[1:] for name, hexv in palette("dark")}
WHITE = "FFFFFF"
REF_PT, REF_W, REF_H, EDGE = 24, 900, 220, 28
# Title and slide-number boxes, shared by the master and the layouts
TITLE_BOX, NUMBER_BOX = (64, 40, 1792, 90), (16, 10, 120, 24)
# Picture formats python-pptx places
FIGURE_FORMATS = {"PNG", "JPEG", "GIF", "BMP", "TIFF"}
REF_PROMPT = [[("Author+ YY ", False), ("Journal", True)]]
# Figure boxes of the layouts, in design points: (x, y, width, height)
BOX = {
    ("Figure", 10): (640, 190, 1216, 790),
    ("Figure, large", 10): (160, 185, 1600, 790),
    ("Two figures", 10): (64, 200, 880, 770),
    ("Two figures", 11): (976, 200, 880, 770),
}


def _attr(text):
    """Escape text for an XML attribute value."""
    return escape(text, {'"': "&quot;"})


class _Geo:
    """Convert points of the 1920 x 1080 design to EMU and font units at one scale."""

    def __init__(self, scale):
        self.s = scale

    def e(self, pt):
        return int(round(pt * self.s * 12700))

    def sz(self, pt):  # hundredths of a point
        return int(round(pt * self.s * 100))


def _rpr(g, size, color, face="Geist", italic=False, spc=0, tag="a:defRPr"):
    attrs = f'lang="en-GB" sz="{g.sz(size)}" b="0" i="{int(italic)}"'
    if spc:
        attrs += f' spc="{g.sz(spc / 100)}"'
    return (
        f'<{tag} {attrs}><a:solidFill><a:srgbClr val="{color}"/></a:solidFill>'
        f'<a:latin typeface="{face}"/><a:ea typeface="{face}"/><a:cs typeface="{face}"/></{tag}>'
    )


class _Layout:
    """Collect the shapes of one layout as XML."""

    def __init__(self, g, name, ground):
        self.g, self.name, self.ground = g, name, ground
        self.shapes, self.images, self._next = [], {}, 2

    def _id(self):
        self._next += 1
        return self._next

    def _xfrm(self, x, y, w, h):
        g = self.g
        return (
            f'<a:xfrm><a:off x="{g.e(x)}" y="{g.e(y)}"/>'
            f'<a:ext cx="{g.e(w)}" cy="{g.e(h)}"/></a:xfrm>'
        )

    def ph(
        self,
        kind,
        idx,
        name,
        box,
        size,
        color,
        prompt,
        face="Geist",
        italic=False,
        align="l",
        anchor="t",
        spc=0,
        lines=None,
    ):
        """Add a placeholder; ``lines`` holds the prompt as lists of (text, italic) runs."""
        typ = {
            "title": 'type="title"',
            "body": f'type="body" idx="{idx}"',
            "figure": f'idx="{idx}"',
            "number": f'type="sldNum" sz="quarter" idx="{idx}"',
        }[kind]
        lvl = (
            f'<a:lvl1pPr marL="0" indent="0" algn="{align}"><a:lnSpc><a:spcPct val="100000"/>'
            '</a:lnSpc><a:spcBef><a:spcPts val="0"/></a:spcBef><a:buNone/>'
            + _rpr(self.g, size, color, face, italic, spc)
            + "</a:lvl1pPr>"
        )
        runs = lines if lines is not None else [[(prompt, italic)]]
        paras = "".join(
            "<a:p>"
            + "".join(
                f'<a:r><a:rPr lang="en-GB" i="{int(it)}" dirty="0"/><a:t>{escape(t)}</a:t></a:r>'
                for t, it in line
            )
            + "</a:p>"
            for line in runs
        )
        if kind == "number":
            paras = (
                '<a:p><a:fld id="{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}" type="slidenum">'
                '<a:rPr lang="en-GB"/><a:t>‹#›</a:t></a:fld></a:p>'
            )
        body = (
            '<p:txBody><a:bodyPr vert="horz" wrap="square" lIns="0" tIns="0" rIns="0" bIns="0" '
            f'anchor="{anchor}"><a:noAutofit/></a:bodyPr><a:lstStyle>{lvl}</a:lstStyle>'
            f"{paras}</p:txBody>"
        )
        self.shapes.append(
            f'<p:sp><p:nvSpPr><p:cNvPr id="{self._id()}" name="{_attr(name)}"/><p:cNvSpPr>'
            f'<a:spLocks noGrp="1"/></p:cNvSpPr><p:nvPr><p:ph {typ}/></p:nvPr></p:nvSpPr>'
            f"<p:spPr>{self._xfrm(*box)}</p:spPr>{body}</p:sp>"
        )

    def text(self, name, box, s, size, color, align="l"):
        """Add fixed text, not a placeholder."""
        self.shapes.append(
            f'<p:sp><p:nvSpPr><p:cNvPr id="{self._id()}" name="{_attr(name)}"/>'
            f'<p:cNvSpPr txBox="1"/><p:nvPr userDrawn="1"/></p:nvSpPr><p:spPr>{self._xfrm(*box)}'
            '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:noFill/></p:spPr>'
            '<p:txBody><a:bodyPr wrap="none" lIns="0" tIns="0" rIns="0" bIns="0" anchor="ctr">'
            f'<a:noAutofit/></a:bodyPr><a:lstStyle/><a:p><a:pPr algn="{align}"/><a:r>'
            + _rpr(self.g, size, color, tag="a:rPr")
            + f"<a:t>{escape(s)}</a:t></a:r></a:p></p:txBody></p:sp>"
        )

    def image(self, name, path, box):
        key = f"__IMG{len(self.images)}__"
        self.images[key] = path
        self.shapes.append(
            f'<p:pic><p:nvPicPr><p:cNvPr id="{self._id()}" name="{_attr(name)}"/><p:cNvPicPr>'
            '<a:picLocks noChangeAspect="1"/></p:cNvPicPr><p:nvPr userDrawn="1"/></p:nvPicPr>'
            f'<p:blipFill><a:blip r:embed="{key}"/><a:stretch><a:fillRect/></a:stretch>'
            f'</p:blipFill><p:spPr>{self._xfrm(*box)}<a:prstGeom prst="rect"><a:avLst/>'
            "</a:prstGeom></p:spPr></p:pic>"
        )

    def xml(self):
        bg = (
            f'<p:bg><p:bgPr><a:solidFill><a:srgbClr val="{self.ground}"/></a:solidFill>'
            "<a:effectLst/></p:bgPr></p:bg>"
        )
        return (
            f'<p:sldLayout {_NS} preserve="1" userDrawn="1"><p:cSld name="{_attr(self.name)}">'
            f'{bg}<p:spTree><p:nvGrpSpPr><p:cNvPr id="1" name=""/><p:cNvGrpSpPr/><p:nvPr/>'
            '</p:nvGrpSpPr><p:grpSpPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="0" cy="0"/>'
            '<a:chOff x="0" y="0"/><a:chExt cx="0" cy="0"/></a:xfrm></p:grpSpPr>'
            + "".join(self.shapes)
            + "</p:spTree></p:cSld><p:clrMapOvr><a:masterClrMapping/></p:clrMapOvr></p:sldLayout>"
        )


def _layouts(g, figures):
    """Return the layouts in the order of :data:`LAYOUTS`."""
    out = []

    def new(name, ground=INK["950"]):
        lay = _Layout(g, name, ground)
        out.append(lay)
        return lay

    def refs(lay):
        y = 1080 - EDGE - REF_H
        lay.ph(
            "body",
            20,
            "Reference left",
            (EDGE, y, REF_W, REF_H),
            REF_PT,
            INK["300"],
            "",
            anchor="b",
            lines=REF_PROMPT,
        )
        lay.ph(
            "body",
            21,
            "Reference right",
            (1920 - EDGE - REF_W, y, REF_W, REF_H),
            REF_PT,
            INK["300"],
            "",
            align="r",
            anchor="b",
            lines=REF_PROMPT,
        )

    def number(lay):
        lay.ph("number", 30, "Slide number", NUMBER_BOX, 16, INK["400"], "")

    def logo(lay, x=1720, y=930, h=100):
        w = h * 36.75 / 20.25
        lay.image("Forming Worlds logo", assets_dir() / "fwl-logo-white.png", (x, y, w, h))
        lay.text(
            "URL", (x - 40, y + h + 4, w + 40, 22), "formingworlds.space", 17, INK["100"], align="r"
        )

    def title(lay):
        lay.ph(
            "title",
            0,
            "Title",
            TITLE_BOX,
            72,
            WHITE,
            "Title states the finding in one line",
            face="Geist Light",
        )

    def figure(lay, idx, name):
        if figures:
            lay.ph(
                "figure",
                idx,
                name,
                BOX[(lay.name, idx)],
                24,
                INK["400"],
                "Figure",
                align="ctr",
                anchor="ctr",
            )

    lay = new("Title")
    lay.ph("title", 0, "Title", (64, 70, 1600, 300), 120, WHITE, "Talk title", face="Geist Light")
    lay.ph(
        "body",
        1,
        "Speaker",
        (1000, 330, 856, 50),
        30,
        WHITE,
        "SPEAKER NAME",
        face="Geist Medium",
        align="r",
        spc=600,
    )
    lay.ph(
        "body",
        2,
        "Key papers",
        (64, 470, 760, 300),
        22,
        INK["100"],
        "Optional: key papers, one per line",
    )
    lay.ph(
        "body",
        3,
        "Venue and date",
        (EDGE, 1080 - 120 - EDGE, 520, 120),
        22,
        INK["100"],
        "",
        anchor="b",
        lines=[[("Venue, city", False)], [("Date", False)]],
    )
    logo(lay)

    lay = new("Question")
    lay.ph(
        "title",
        0,
        "Question",
        (80, 120, 1760, 260),
        88,
        WHITE,
        "How do rocky worlds form and evolve?",
        face="Geist Light",
        align="ctr",
        anchor="b",
    )
    lay.ph(
        "body",
        1,
        "Sub-question",
        (260, 400, 1400, 160),
        44,
        INK["100"],
        "Sub-question for the next build",
        face="Geist Light",
        align="ctr",
    )
    number(lay)

    lay = new("Section")
    lay.ph(
        "body",
        1,
        "Section label",
        (64, 380, 1200, 40),
        24,
        MAGMA400,
        "SECTION",
        face="Geist Mono",
        spc=300,
    )
    lay.ph(
        "title",
        0,
        "Section title",
        (64, 430, 1500, 260),
        96,
        WHITE,
        "Section statement",
        face="Geist Light",
    )
    number(lay)

    lay = new("Figure")
    title(lay)
    figure(lay, 10, "Figure")
    lay.ph(
        "body",
        1,
        "Side text",
        (64, 400, 520, 420),
        34,
        INK["100"],
        "Short line on what the figure shows",
        face="Geist Light",
    )
    refs(lay)
    number(lay)

    lay = new("Figure, large")
    title(lay)
    figure(lay, 10, "Figure")
    refs(lay)
    number(lay)

    lay = new("Two figures")
    title(lay)
    figure(lay, 10, "Figure left")
    figure(lay, 11, "Figure right")
    refs(lay)
    number(lay)

    lay = new("Summary")
    for i in range(4):
        y = 90 + i * 225
        lay.ph(
            "body",
            1 + i,
            f"Point {i + 1}",
            (64, y, 1240, 200),
            44,
            WHITE,
            f"{i + 1}. Take-home point",
            face="Geist Medium",
        )
        lay.ph(
            "body",
            11 + i,
            f"Point {i + 1} references",
            (1350, y + 8, 500, 200),
            REF_PT,
            INK["300"],
            "",
            lines=[[("Author+ YY", False)], [("Author & Author YY", False)]],
        )
    logo(lay)

    lay = new("Statement", INK["980"])
    lay.ph(
        "title",
        0,
        "Statement",
        (200, 330, 1520, 420),
        64,
        WHITE,
        "A common assumption, stated plainly.",
        face="Geist Light",
        italic=True,
        align="ctr",
        anchor="ctr",
    )
    number(lay)

    lay = new("Team")
    title(lay)
    refs(lay)
    number(lay)

    lay = new("Closing")
    lay.ph(
        "title",
        0,
        "Closing statement",
        (160, 120, 1600, 260),
        88,
        WHITE,
        "Closing question or statement",
        face="Geist Light",
        align="ctr",
        anchor="b",
    )
    lay.ph(
        "body",
        1,
        "Announcement",
        (360, 400, 1200, 120),
        32,
        INK["100"],
        "Open positions, meetings",
        face="Geist Light",
        align="ctr",
    )
    logo(lay)
    return out


def _theme_xml(blob):
    """Return the theme part with the lab colour and font schemes."""
    from lxml import etree

    t = etree.fromstring(blob)
    cs = t.find(f".//{_A}clrScheme")
    cs.set("name", "Forming Worlds")
    values = {
        "dk1": INK["950"],
        "lt1": WHITE,
        "dk2": INK["800"],
        "lt2": INK["100"],
        "accent1": MAGMA400,
        "accent2": SIGNAL400,
        "accent3": DARK["fern"],
        "accent4": DARK["ochre"],
        "accent5": DARK["iris"],
        "accent6": DARK["garnet"],
        "hlink": SIGNAL400,
        "folHlink": SIGNAL300,
    }
    for key, hexv in values.items():
        el = cs.find(f"{_A}{key}")
        for child in list(el):
            el.remove(child)
        etree.SubElement(el, f"{_A}srgbClr", val=hexv)
    fs = t.find(f".//{_A}fontScheme")
    fs.set("name", "Forming Worlds")
    for tag, face in (("majorFont", "Geist Light"), ("minorFont", "Geist")):
        fs.find(f"{_A}{tag}/{_A}latin").set("typeface", face)
    return etree.tostring(t, xml_declaration=True, encoding="UTF-8", standalone=True)


def _fill(shape, lines):
    """Write runs into a placeholder: style "" plain, "i" italic, "#RRGGBB" colour."""
    from pptx.dml.color import RGBColor

    tf = shape.text_frame
    for i, runs in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        for text, style in runs:
            r = p.add_run()
            r.text = text
            if style == "i":
                r.font.italic = True
            elif style.startswith("#"):
                r.font.color.rgb = RGBColor.from_string(style[1:])


def _samples(prs, g, figure):
    """Add one sample slide per layout; ``figure`` fills the figure boxes when given."""
    from PIL import Image
    from pptx.util import Emu

    by = {lay.name: lay for lay in prs.slide_layouts}
    hl = "#" + MAGMA400

    def slide(name, fill, figs=()):
        s = prs.slides.add_slide(by[name])
        phs = {ph.placeholder_format.idx: ph for ph in s.placeholders}
        for idx, lines in fill.items():
            _fill(phs[idx], lines)
        for idx in figs if figure is not None else ():
            if idx in phs:  # the sample figure takes the place of the empty box
                phs[idx].element.getparent().remove(phs[idx].element)
            x, y, w, h = BOX[(name, idx)]
            with Image.open(figure) as im:
                aspect = im.width / im.height
            fw, fh = (w, w / aspect) if w / aspect <= h else (h * aspect, h)
            s.shapes.add_picture(
                str(figure),
                Emu(g.e(x + (w - fw) / 2)),
                Emu(g.e(y + (h - fh) / 2)),
                Emu(g.e(fw)),
                Emu(g.e(fh)),
            )

    slide(
        "Title",
        {
            0: [[("Talk title", "")]],
            1: [[("SPEAKER NAME", "")]],
            3: [[("Venue, city", "")], [("Date", "")]],
        },
    )
    slide("Question", {0: [[("How do ", ""), ("rocky worlds", hl), (" form and evolve?", "")]]})
    slide(
        "Section",
        {1: [[("PART 2", "")]], 0: [[("Lava worlds open a window into the interior", "")]]},
    )
    slide(
        "Figure",
        {
            0: [[("Lava world atmospheres are fed from the ", ""), ("interior", hl)]],
            1: [[("Outgassing follows the redox state of the magma", "")]],
            20: [[("Nicholls+ 25 ", ""), ("MNRAS", "i")]],
        },
        figs=[10],
    )
    slide(
        "Figure, large",
        {
            0: [[("Residual atmospheres are mass-sorted", "")]],
            21: [[("Lichtenberg+ 25 ", ""), ("Science", "i")]],
        },
        figs=[10],
    )
    slide(
        "Two figures",
        {
            0: [[("Sulfur species trace the interior phase state", "")]],
            20: [[("Teske+ 25 ", ""), ("ApJL", "i")]],
            21: [[("Lichtenberg & Miguel 25 ", ""), ("Treatise on Geochemistry", "i")]],
        },
        figs=[10, 11],
    )
    slide(
        "Summary",
        {
            1: [[("1. First take-home point", "")]],
            11: [[("Author+ 25", "")], [("Author & Author 26", "")]],
            2: [[("2. Second take-home point", "")]],
            12: [[("Author+ 24", "")]],
            3: [[("3. Third take-home point", "")]],
            13: [[("Author+ 26a,b", "")]],
        },
    )
    slide("Statement", {0: [[("A common assumption, stated plainly.", "")]]})
    slide("Team", {0: [[("Forming Worlds Lab", "")]]})
    slide(
        "Closing",
        {
            0: [[("How do ", ""), ("rocky worlds", hl), (" form and evolve?", "")]],
            1: [[("Open positions: formingworlds.space", "")]],
        },
    )


def _master(master, g):
    """Set the master to the dark colour map, 16:9 placeholders and Geist text styles."""
    from pptx.oxml import parse_xml
    from pptx.util import Emu

    el = master.element
    p_ns = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
    bg = el.find(f"{p_ns}cSld/{p_ns}bg")
    bg.getparent().replace(
        bg,
        parse_xml(
            f'<p:bg {_NS}><p:bgPr><a:solidFill><a:srgbClr val="{INK["950"]}"/></a:solidFill>'
            "<a:effectLst/></p:bgPr></p:bg>"
        ),
    )
    # Text and background 1 map to the light and dark theme colours: text that takes no
    # colour of its own (a new text box, a table, list levels 2 and below) is light.
    clr = el.find(f"{p_ns}clrMap")
    clr.attrib.update({"bg1": "dk1", "tx1": "lt1", "bg2": "dk2", "tx2": "lt2"})
    boxes = {
        "TITLE": TITLE_BOX,
        "BODY": (64, 190, 1792, 790),
        "DATE": (1660, 10, 232, 24),
        "FOOTER": (1200, 10, 440, 24),
        "SLIDE_NUMBER": NUMBER_BOX,
    }
    # a placeholder type this design does not place keeps its position
    for ph in (p for p in master.placeholders if p.placeholder_format.type.name in boxes):
        x, y, w, h = boxes[ph.placeholder_format.type.name]
        ph.left, ph.top, ph.width, ph.height = (Emu(g.e(v)) for v in (x, y, w, h))

    def level(n, size, color, face, bullet):
        mar = 0 if not bullet else 40 * n
        bu = f'<a:buFont typeface="{face}"/><a:buChar char="•"/>' if bullet else "<a:buNone/>"
        return (
            f'<a:lvl{n}pPr marL="{g.e(mar)}" indent="{-g.e(30) if bullet else 0}">'
            f'<a:spcBef><a:spcPts val="{g.sz(8)}"/></a:spcBef>{bu}'
            + _rpr(g, size, color, face)
            + f"</a:lvl{n}pPr>"
        )

    styles = (
        f"<p:txStyles {_NS}><p:titleStyle>{level(1, 72, WHITE, 'Geist Light', False)}"
        "</p:titleStyle>"
        "<p:bodyStyle>"
        + "".join(level(n, 34 - 4 * min(n - 1, 3), INK["100"], "Geist", True) for n in range(1, 6))
        + f"</p:bodyStyle><p:otherStyle>{level(1, 24, INK['100'], 'Geist', False)}</p:otherStyle>"
        "</p:txStyles>"
    )
    old = el.find(f"{p_ns}txStyles")
    old.getparent().replace(old, parse_xml(styles))


def build_slides(out, keynote=False, figure=None):
    """Write the slide master file.

    Parameters
    ----------
    out : str or pathlib.Path
        Path of the ``.pptx`` to write; missing parent folders are created.
    keynote : bool, optional
        Write the Keynote source: 1920 x 1080 pt and no figure placeholders.
        Otherwise the file is 13.333 x 7.5 in with figure placeholders.
    figure : str or pathlib.Path, optional
        Image placed on the sample figure slides.

    Returns
    -------
    pathlib.Path
        Path of the written file.

    Raises
    ------
    ImportError
        If ``python-pptx`` is not installed; the message names the ``slides`` extra.
    ValueError
        If ``figure`` is not a PNG, JPEG, GIF, BMP or TIFF file that can be read.
    RuntimeError
        If the python-pptx base template has fewer layouts than :data:`LAYOUTS`.
    """
    try:
        from pptx import Presentation
        from pptx.oxml import parse_xml
        from pptx.util import Emu
    except ImportError as err:
        raise ImportError(
            'the slide masters need python-pptx: pip install "fwl-visual-language[slides]"'
        ) from err
    if figure is not None:
        from PIL import Image

        try:
            with Image.open(figure) as im:
                kind = im.format
                im.load()  # decodes the whole image; verify() skips most formats
        except Exception as err:  # Pillow raises OSError, SyntaxError, ValueError, ...
            raise ValueError(f"cannot read the figure {figure}: {err}") from err
        if kind not in FIGURE_FORMATS:
            raise ValueError(f"the figure {figure} is {kind}; use one of {sorted(FIGURE_FORMATS)}")

    g = _Geo(1.0 if keynote else 0.5)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(g.e(1920)), Emu(g.e(1080))
    master = prs.slide_master
    _master(master, g)
    theme = next(r.target_part for r in master.part.rels.values() if r.reltype.endswith("/theme"))
    theme.blob = _theme_xml(theme.blob)
    props = prs.core_properties
    props.title, props.author, props.last_modified_by = (
        "Forming Worlds Lab slide masters",
        "Forming Worlds Lab",
        "Forming Worlds Lab",
    )
    props.created = props.modified = datetime(2026, 10, 18)
    props.revision = 1

    defs = _layouts(g, figures=not keynote)
    while len(prs.slide_layouts) > len(defs):  # the base template has 11 layouts
        prs.slide_layouts.remove(prs.slide_layouts[-1])
    layouts = list(prs.slide_layouts)
    if len(layouts) != len(defs):
        raise RuntimeError(f"expected {len(defs)} layouts in the base template, got {len(layouts)}")
    for lay, layout in zip(defs, layouts, strict=True):
        xml = lay.xml()
        for key, path in lay.images.items():
            _, rid = layout.part.get_or_add_image_part(str(path))
            xml = xml.replace(f'r:embed="{key}"', f'r:embed="{rid}"')
        # Replace the content in place, so every wrapper of this layout sees the new XML.
        old, new = layout.element, parse_xml(xml)
        old.attrib.clear()
        old.attrib.update(new.attrib)
        for child in list(old):
            old.remove(child)
        for child in list(new):
            old.append(child)
    _samples(prs, g, figure)
    out = Path(out)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    return out


def main(argv=None):
    """Command line entry point: ``python -m fwl_visual_language.slides OUT``."""
    parser = argparse.ArgumentParser(
        prog="python -m fwl_visual_language.slides",
        description="Write the Forming Worlds Lab slide masters as a .pptx.",
    )
    parser.add_argument("out", help="path of the .pptx to write")
    parser.add_argument(
        "--keynote",
        action="store_true",
        help="write the 1920 x 1080 Keynote source without figure placeholders",
    )
    parser.add_argument("--figure", help="image for the sample figure slides")
    args = parser.parse_args(argv)
    print(build_slides(args.out, args.keynote, args.figure))


if __name__ == "__main__":
    main()
